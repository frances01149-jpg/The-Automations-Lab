#!/usr/bin/env python3
"""
Operations Hub workbook generator.

Builds a formula-driven Excel workbook with two independent systems:

  SYSTEM 1  Inventory Management   (Control Panel, Inventory Master, Weekly Count,
                                    Order List, Inventory Dashboard)
  SYSTEM 2  Pantry Pilot Tracker   (Pilot Products, Restock Log, Feedback Log,
                                    Pilot Dashboard)

The product master comes from a source workbook with the columns
Category | Item Name | Item # | Recent Date Received | PAR Value
(English or Spanish headers are recognised). Every other number in the
workbook is a formula, so the same file works for any office: swap the
source data and the pilot dates and the whole system recalculates.

Usage
  python build_ops_hub.py --source "SJUB - Inventario.xlsx" --site SJUB
  python build_ops_hub.py --site SJUB                    # empty template
  python build_ops_hub.py --demo                         # sample data
  python build_ops_hub.py --site SJUB --xltx             # Excel template

See ops-hub/README.md for the full options.
"""
from __future__ import annotations

import argparse
import datetime as dt
import math
import random
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as col_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

# --------------------------------------------------------------------------- #
# Capacity (rows pre-built with formulas; blank rows stay invisible)
# --------------------------------------------------------------------------- #
CAP_PRODUCTS = 300      # source / master / order list rows
CAP_COUNTS = 3000       # weekly count log rows
CAP_REFILLS = 1500      # restock log rows
CAP_FEEDBACK = 1000     # feedback log rows
CAP_PILOT = 80          # pilot product rows
CAP_CATS = 40           # distinct categories
CAP_LIST = 30           # staff / team / category-map list rows
TOP_N = 10              # rows in dashboard "top" charts

HDR = 5                 # header row of every table
FIRST = HDR + 1         # first data row

# --------------------------------------------------------------------------- #
# Palette
# --------------------------------------------------------------------------- #
INK = "0B1A3A"          # title bar
NAVY = "142A5E"         # table headers, primary cards
NAV_BG = "1D3570"       # navigation strip
NAV_ON = "DCE566"       # active nav item (lime)
SLATE = "5B6B8C"
MUTED = "8A96AE"
PAGE = "FFFFFF"
BAND = "F3F5FA"
LINE = "D9DFEC"
PANEL = "EEF2FA"

IN_FILL = "DCE9FF"      # manual entry cells (blue)
IN_FONT = "1446A0"
IN_HDR = "2458C6"
SET_FILL = "FFF6D6"     # one-time setup cells (yellow)
SET_FONT = "6B4E00"
SET_HDR = "B08900"

OK_BG, OK_FG = "D8F3E4", "11693A"
LOW_BG, LOW_FG = "FFEBC7", "8A5300"
CRIT_BG, CRIT_FG = "FBD5D5", "A11A1A"
NONE_BG, NONE_FG = "E7EAF1", "5B6B8C"

C_GREEN, C_AMBER, C_RED = "1E8E5A", "E0981A", "C8373A"
C_TEAL, C_PURPLE, C_NAVY, C_BLUE = "12808A", "5B4BB7", "142A5E", "2458C6"

FONT = "Segoe UI"

# --------------------------------------------------------------------------- #
# Sheet names
# --------------------------------------------------------------------------- #
S_HOME = "Home"
S_CP = "Control Panel"
S_SRC = "Source Data"
S_MS = "Inventory Master"
S_CL = "Weekly Count"
S_OL = "Order List"
S_ID = "Inventory Dashboard"
S_PP = "Pilot Products"
S_RL = "Restock Log"
S_FB = "Feedback Log"
S_PD = "Pilot Dashboard"
S_EN = "_Engine"

NAV = [
    (S_HOME, "⌂  HOME"),
    (S_CP, "⚙  CONTROL"),
    (S_MS, "INVENTORY"),
    (S_CL, "✎ COUNT"),
    (S_OL, "ORDERS"),
    (S_ID, "◧ INV. DASH"),
    (S_PP, "PILOT ITEMS"),
    (S_RL, "✎ RESTOCK"),
    (S_FB, "✎ FEEDBACK"),
    (S_PD, "◧ PILOT DASH"),
]

TAB = {
    S_HOME: "DCE566", S_CP: "5B6B8C", S_SRC: "8A96AE",
    S_MS: "142A5E", S_CL: "2458C6", S_OL: "142A5E", S_ID: "0B1A3A",
    S_PP: "12808A", S_RL: "2458C6", S_FB: "2458C6", S_PD: "0B5E66",
}

# Number formats (stored locale-independently in the file)
F_DATE = "mmm d, yyyy"
F_MONTH = "mmm yyyy"
F_TIME = "h:mm AM/PM"
F_STAMP = "mmm d, h:mm AM/PM"
F_INT = "#,##0"
F_DEC = "0.0"
F_PCT = "0%"


# --------------------------------------------------------------------------- #
# Style helpers
# --------------------------------------------------------------------------- #
def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def font(size=10, bold=False, color="1F2A44", italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic,
                underline=underline)


def side(color=LINE, style="thin"):
    return Side(style=style, color=color)


BOTTOM = Border(bottom=side())
BOX = Border(left=side(), right=side(), top=side(), bottom=side())

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
RIGHT = Alignment(horizontal="right", vertical="center", indent=1)
UNLOCK = Protection(locked=False)


def style(cell, *, f=None, fl=None, al=None, bd=None, nf=None, unlock=False):
    if f is not None:
        cell.font = f
    if fl is not None:
        cell.fill = fl
    if al is not None:
        cell.alignment = al
    if bd is not None:
        cell.border = bd
    if nf is not None:
        cell.number_format = nf
    if unlock:
        cell.protection = UNLOCK
    return cell


def merge(ws, r1, c1, r2, c2, value=None, **kw):
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            if "fl" in kw:
                cell.fill = kw["fl"]
            if "bd" in kw:
                cell.border = kw["bd"]
    cell = ws.cell(r1, c1)
    if value is not None:
        cell.value = value
    style(cell, **kw)
    return cell


def link(cell, sheet, tip=None):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1",
                               tooltip=tip or f"Open {sheet}", display=str(cell.value))


def rng(sheet, col, r1, r2):
    return f"'{sheet}'!${col}${r1}:${col}${r2}"


def cell_ref(sheet, coord):
    c, r = re.match(r"([A-Z]+)(\d+)", coord).groups()
    return f"'{sheet}'!${c}${r}"


# --------------------------------------------------------------------------- #
# Source loading
# --------------------------------------------------------------------------- #
HEADER_KEYS = {
    "category": ["category", "categoría", "categoria", "cat"],
    "name": ["item name", "nombre", "producto", "product", "description",
             "descripción", "descripcion", "artículo", "articulo", "item"],
    "number": ["item #", "item#", "item no", "item number", "no.", "número",
               "numero", "código", "codigo", "sku", "code", "#"],
    "received": ["recent date received", "date received", "received",
                 "recibido", "fecha"],
    "par": ["par value", "par level", "par", "nivel par"],
    "expiration": ["expiration", "expiry", "exp. date", "exp date", "vencimiento",
                   "caducidad"],
}


def _match(text, keys):
    t = str(text).strip().lower()
    return any(t == k or t.startswith(k) or k in t for k in keys)


def load_source(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    best = None
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for hr, row in enumerate(rows[:25]):
            cols = {}
            for ci, v in enumerate(row):
                if v is None:
                    continue
                # most specific keys first so "Item #" is not taken as "name"
                for key in ("par", "expiration", "received", "number", "category", "name"):
                    if key not in cols and _match(v, HEADER_KEYS[key]):
                        if key == "name" and _match(v, HEADER_KEYS["number"]) \
                                and str(v).strip().lower() != "item":
                            continue
                        cols[key] = ci
                        break
            if {"name", "par"} <= cols.keys():
                best = (ws.title, rows, hr, cols)
                break
        if best:
            break
    if not best:
        sys.exit(f"Could not find a header row with Item Name and PAR in {path}")
    title, rows, hr, cols = best
    out = []
    section = ""
    for row in rows[hr + 1:]:
        get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None
        name = get("name")
        if name is None or str(name).strip() == "":
            # a row with text but no item name is a section banner, e.g. "Snacks (...)"
            texts = [str(v).strip() for v in row if v not in (None, "") and str(v).strip()]
            if len(texts) == 1:
                section = texts[0]
            continue
        rec = get("received")
        if isinstance(rec, dt.datetime):
            rec = rec.date()
        par = get("par")
        try:
            par = float(par) if par not in (None, "") else None
            if par is not None and par.is_integer():
                par = int(par)
        except (TypeError, ValueError):
            par = None
        num = get("number")
        if isinstance(num, float) and num.is_integer():
            num = int(num)
        if num is not None and str(num).strip() == "":
            num = None
        exp = get("expiration")
        if isinstance(exp, dt.datetime):
            exp = exp.date()
        elif isinstance(exp, str):
            exp = exp.strip() or None
        out.append({
            "category": (str(get("category")).strip() if get("category") else "Uncategorized"),
            "name": str(name).strip(),
            "number": num,
            "received": rec if isinstance(rec, (dt.date, dt.datetime)) else None,
            "par": par,
            "expiration": exp,
            "section": section,
        })
    print(f"Loaded {len(out)} products from '{path.name}' / sheet '{title}' "
          f"(header row {hr + 1}; columns: {', '.join(sorted(cols))})")
    return out


SNACK_WORDS = ["snack", "botana", "galleta", "cookie", "biscuit", "chip", "candy",
               "dulce", "granola", "nut", "cracker", "pretzel", "popcorn", "fruit",
               "fruta", "mint"]
BEV_WORDS = ["beverage", "bebida", "drink", "soda", "water", "agua", "juice",
             "jugo", "coffee", "café", "cafe", "tea", "té", "refresco", "leche"]


NOT_BEV_WORDS = ["lid", "cup", "stir", "sweet", "straw"]   # drink supplies, not drinks


def guess_pilot_categories(products):
    """Snack = categories under a 'Snack' section banner or with snack words;
    Beverage = categories with drink words (supplies such as lids and cups excluded)."""
    snack, bev = [], []
    for p in products:
        c, low = p["category"], p["category"].lower()
        if c in snack or c in bev:
            continue
        is_bev = any(w in low for w in BEV_WORDS) and not any(w in low for w in NOT_BEV_WORDS)
        if is_bev:
            bev.append(c)
        elif "snack" in p.get("section", "").lower() or any(w in low for w in SNACK_WORDS):
            snack.append(c)
    return snack, bev


# --------------------------------------------------------------------------- #
# Demo data (clearly labelled sample content for trying the system)
# --------------------------------------------------------------------------- #
def demo_products():
    catalog = {
        "Snacks": ["Potato Chips 1 oz", "Pretzel Bites", "Granola Bar", "Trail Mix",
                   "Cheese Crackers", "Fruit Snacks", "Popcorn Pack", "Protein Bar",
                   "Almonds 1.5 oz", "Chocolate Cookies"],
        "Beverages": ["Bottled Water 16 oz", "Sparkling Water", "Cola 12 oz",
                      "Diet Cola 12 oz", "Orange Juice", "Iced Tea", "Sports Drink",
                      "Cold Brew Can"],
        "Coffee & Tea": ["Ground Coffee 2 lb", "Decaf Coffee Pods", "Green Tea Bags",
                         "Creamer Singles", "Sugar Packets"],
        "Paper Goods": ["Paper Towels", "Napkins", "Paper Plates 9 in", "Hot Cups 12 oz",
                        "Cold Cups 16 oz"],
        "Cleaning Supplies": ["Disinfectant Wipes", "Dish Soap", "Trash Bags 33 gal",
                              "Glass Cleaner", "Hand Soap Refill"],
        "Kitchen Supplies": ["Plastic Forks", "Plastic Spoons", "Stir Sticks",
                             "Aluminum Foil", "Food Storage Bags"],
    }
    rnd = random.Random(7)
    out, n = [], 1001
    for cat, names in catalog.items():
        for name in names:
            out.append({
                "category": cat, "name": name, "number": n,
                "received": dt.date(2026, 9, rnd.randint(1, 20)),
                "par": rnd.choice([12, 18, 24, 24, 36, 48, 60]),
                "expiration": dt.date(2026, rnd.randint(10, 12), rnd.randint(1, 28))
                if cat in ("Snacks", "Beverages") else "N/A",
            })
            n += 1
    return out


def demo_entries(products, pilot_start, snack_cats, bev_cats):
    rnd = random.Random(42)
    staff = ["A. Rivera", "M. Chen", "J. Ortiz", "K. Patel"]
    counts = []
    level = {p["name"]: p["par"] * rnd.uniform(0.6, 1.0) for p in products}
    for w in range(4):
        d = pilot_start + dt.timedelta(days=7 * w)
        who = staff[w % len(staff)]
        for p in products:
            use = p["par"] * rnd.uniform(0.05, 0.3)
            if level[p["name"]] - use < p["par"] * 0.15 and rnd.random() < 0.6:
                level[p["name"]] = p["par"] * rnd.uniform(0.8, 1.0)   # restocked
            else:
                level[p["name"]] = max(0, level[p["name"]] - use)
            counts.append((d, p["name"], round(level[p["name"]]), who))

    pilot = [p for p in products if p["category"] in snack_cats + bev_cats][:14]
    refills = []
    for p in pilot:
        t = dt.datetime.combine(pilot_start, dt.time(8, 30))
        qty_before = 0
        target = max(6, p["par"] // 2)
        while t.date() <= pilot_start + dt.timedelta(days=22):
            added = target - qty_before
            note = "Initial placement" if qty_before == 0 and not any(
                r[2] == p["name"] for r in refills) else rnd.choice(
                ["", "", "", "Shelf empty on arrival" if qty_before == 0 else "",
                 "High demand after lunch", "Rotated stock (FIFO)"])
            refills.append((t.date(), t.time(), p["name"], qty_before, added,
                            rnd.choice(staff), note))
            gap = rnd.uniform(1.5, 4.5)
            t = t + dt.timedelta(days=gap)
            t = t.replace(hour=rnd.choice([8, 9, 11, 13, 15]), minute=rnd.choice([0, 15, 30, 45]),
                          second=0, microsecond=0)
            qty_before = max(0, int(target * rnd.uniform(-0.1, 0.35)))
    refills.sort(key=lambda r: (r[0], r[1]))

    comments = {
        "Positive": ["Love having drinks close by", "Great snack variety",
                     "Saves a trip downstairs", "Team morale boost"],
        "Neutral": ["Fine as is", "Would use it more with healthier options"],
        "Concern": ["Water runs out by 2 PM", "Shelf was empty Friday afternoon"],
        "Suggestion": ["Add sugar-free options", "Add oat milk for coffee",
                       "Please restock earlier on Mondays"],
    }
    teams = ["Front Office", "Operations", "Finance", "IT", "HR"]
    fb = []
    for i in range(28):
        d = pilot_start + dt.timedelta(days=rnd.randint(0, 22))
        typ = rnd.choices(list(comments), weights=[12, 4, 3, 5])[0]
        rating = {"Positive": rnd.choice([4, 5, 5]), "Neutral": 3,
                  "Concern": rnd.choice([2, 3]), "Suggestion": rnd.choice([3, 4])}[typ]
        fu = "Yes" if typ in ("Concern", "Suggestion") else "No"
        status, resolved = ("Closed", None) if fu == "No" else rnd.choice(
            [("Open", None), ("In Review", None), ("Closed", d + dt.timedelta(days=rnd.randint(1, 4)))])
        prod = rnd.choice(["Pantry — General"] + [p["name"] for p in pilot])
        fb.append((d, dt.time(rnd.choice([9, 10, 12, 14, 16]), rnd.choice([0, 20, 40])),
                   typ, rating, prod, rnd.choice(comments[typ]), rnd.choice(teams), fu,
                   rnd.choice(staff) if fu == "Yes" else None, status, resolved))
    fb.sort(key=lambda r: (r[0], r[1]))
    return counts, refills, fb, staff, teams


# --------------------------------------------------------------------------- #
# Builder
# --------------------------------------------------------------------------- #
class Hub:
    def __init__(self, site, products, pilot_start, pilot_days, order_month,
                 snack_cats, bev_cats, as_of=None, demo=None):
        self.site = site
        self.products = products
        self.pilot_start = pilot_start
        self.pilot_days = pilot_days
        self.order_month = order_month
        self.snack_cats = snack_cats
        self.bev_cats = bev_cats
        self.as_of = as_of
        self.demo = demo
        self.wb = Workbook()
        self.ws = {}
        self.names = {}
        self._band_prio = 500
        self.last_col = {}

    # ---- workbook plumbing ------------------------------------------------ #
    def name(self, nm, ref):
        self.wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
        self.names[nm] = ref

    def sheet(self, title, first=False):
        ws = self.wb.active if first else self.wb.create_sheet()
        ws.title = title
        ws.sheet_properties.tabColor = TAB.get(title, "8A96AE")
        ws.sheet_view.showGridLines = False
        ws.sheet_view.showRowColHeaders = False
        ws.sheet_view.zoomScale = 100
        self.ws[title] = ws
        return ws

    def widths(self, ws, widths, hidden=()):
        visible = len(widths) - len(hidden)
        widths = list(widths) + [12] * max(0, len(NAV) + 3 - visible)  # room for the nav bar
        ws.column_dimensions["A"].width = 2.5
        for i, w in enumerate(widths, start=2):
            ws.column_dimensions[col_letter(i)].width = w
            ws.column_dimensions[col_letter(i)].hidden = col_letter(i) in hidden
        last = 2 + len(widths)
        self.last_col[ws.title] = last
        # one breathing column, then hide the rest of the grid (app look)
        ws.column_dimensions[col_letter(last)].width = 2.5
        ws.column_dimensions.group(col_letter(last + 1), "XFD", hidden=True,
                                   outline_level=0)
        return last - 1

    def chrome(self, ws, key, subtitle, last_col, system):
        """Title bar, navigation strip and subtitle shared by every page."""
        ws.row_dimensions[1].height = 36
        ws.row_dimensions[2].height = 26
        ws.row_dimensions[3].height = 30
        ws.row_dimensions[4].height = 8
        for c in range(1, last_col + 2):
            ws.cell(1, c).fill = fill(INK)
            ws.cell(2, c).fill = fill(NAV_BG)
        vis = [c for c in range(2, last_col + 1) if not ws.column_dimensions[col_letter(c)].hidden]
        wid = {c: ws.column_dimensions[col_letter(c)].width or 8.43 for c in vis}
        k, acc = len(vis) - 1, 0.0
        while k > 1 and acc < 24:          # right-hand block wide enough for the date
            acc += wid[vis[k]]
            k -= 1
        split = vis[k]
        merge(ws, 1, 2, 1, split,
              value=f'="◆  "&UPPER(cfg_Site)&"  OPERATIONS HUB   |   {system}"',
              f=font(15, True, "FFFFFF"), fl=fill(INK),
              al=Alignment(vertical="center", indent=0))
        merge(ws, 1, vis[k + 1], 1, vis[-1],
              value="=cfg_AsOf",
              f=font(10, True, NAV_ON), fl=fill(INK),
              al=Alignment(horizontal="right", vertical="center", indent=1),
              nf='"AS OF  "mmm d, yyyy')

        # navigation strip: greedily merge columns until each item is ~12 wide
        total, n, m = sum(wid.values()), len(NAV), len(vis)
        j, acc = 0, 0.0
        for i, (sheet, label) in enumerate(NAV):
            first, own = j, 0.0
            # take columns until this item reaches its share of the total width
            # (and is at least ~9 wide), leaving a column for each remaining item
            while j < m - (n - 1 - i) and (j == first or own < 9
                                           or acc < total * (i + 1) / n - 2):
                acc += wid[vis[j]]
                own += wid[vis[j]]
                j += 1
            if i == n - 1:
                j = m
            start, end = vis[first], vis[j - 1]
            active = sheet == key
            cell = merge(ws, 2, start, 2, end, value=label,
                         f=font(8.5, True, INK if active else "D5DDF0"),
                         fl=fill(NAV_ON if active else NAV_BG), al=CENTER)
            link(cell, sheet)
        merge(ws, 3, 2, 3, last_col, value=subtitle,
              f=font(9.5, False, SLATE, italic=True),
              al=Alignment(vertical="center", wrap_text=True, indent=0))

    def protect(self, ws):
        p = ws.protection
        p.sheet = True
        p.autoFilter = False
        p.sort = False
        p.formatColumns = False
        p.formatRows = False
        p.selectLockedCells = False
        p.selectUnlockedCells = False

    def table(self, ws, name, first_col, last_col, last_row, row_height=20):
        for r in range(FIRST, last_row + 1):
            ws.row_dimensions[r].height = row_height
        ref = f"{col_letter(first_col)}{HDR}:{col_letter(last_col)}{last_row}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False,
                                          showColumnStripes=False)
        ws.add_table(t)

    def headers(self, ws, heads, kinds=None, height=34):
        """kinds: 'calc' (navy), 'in' (blue entry), 'set' (yellow setup)."""
        ws.row_dimensions[HDR].height = height
        for i, h in enumerate(heads):
            k = (kinds or {}).get(h, "calc")
            bg = {"calc": NAVY, "in": IN_HDR, "set": SET_HDR}[k]
            style(ws.cell(HDR, 2 + i, h), f=font(9, True, "FFFFFF"), fl=fill(bg),
                  al=CENTER, bd=Border(bottom=side(INK, "medium")))

    def band(self, ws, key_col, c1, c2, r1, r2):
        """Row banding + separators that only appear on filled rows."""
        area = f"{col_letter(c1)}{r1}:{col_letter(c2)}{r2}"
        k = f"${col_letter(key_col)}{r1}"
        # low priority so status colours always win over banding
        for formula, fl in [(f'AND({k}<>"",MOD(ROW(),2)=0)', fill(BAND)), (f'{k}<>""', None)]:
            self._band_prio += 1
            rule = FormulaRule(formula=[formula], fill=fl, border=Border(bottom=side()))
            rule.priority = self._band_prio
            ws.conditional_formatting.add(area, rule)

    def status_colors(self, ws, area, mapping):
        for text, (bg, fg) in mapping.items():
            ws.conditional_formatting.add(area, CellIsRule(
                operator="equal", formula=[f'"{text}"'], fill=fill(bg),
                font=Font(name=FONT, bold=True, color=fg), stopIfTrue=True))

    def body(self, ws, col, r1, r2, formula=None, nf=None, al=None, kind="calc",
             wrap=False):
        """Fill a column with a row-relative formula (use {r} for the row)."""
        f_in = font(10, False, IN_FONT)
        f_set = font(10, False, SET_FONT)
        f_calc = font(10, False, "1F2A44")
        a = al or (LEFT if wrap else CENTER)
        for r in range(r1, r2 + 1):
            c = ws.cell(r, col)
            if formula is not None:
                c.value = formula.format(r=r, p=r - 1)
            if kind == "in":
                style(c, f=f_in, fl=fill(IN_FILL), al=a, nf=nf, unlock=True,
                      bd=Border(bottom=side("FFFFFF"), right=side("FFFFFF")))
            elif kind == "set":
                style(c, f=f_set, fl=fill(SET_FILL), al=a, nf=nf, unlock=True,
                      bd=Border(bottom=side("FFFFFF"), right=side("FFFFFF")))
            else:
                style(c, f=f_calc, al=a, nf=nf)

    def dv(self, ws, area, kind, formula1=None, formula2=None, prompt=None,
           strict=True, title=None):
        strip = lambda f: f[1:] if isinstance(f, str) and f.startswith("=") else f
        d = DataValidation(type=kind, formula1=strip(formula1), formula2=strip(formula2),
                           allow_blank=True, showErrorMessage=True,
                           errorStyle="stop" if strict else "warning")
        if kind in ("date", "time", "whole", "decimal"):
            d.operator = "between"
        d.error = "Please pick a value from the list or enter a valid value."
        d.errorTitle = "Invalid entry"
        if prompt:
            d.prompt, d.promptTitle, d.showInputMessage = prompt, title or "Entry", True
        ws.add_data_validation(d)
        d.add(area)

    # ---- dashboard widgets -------------------------------------------------- #
    def card(self, ws, r, c1, c2, label, formula, nf, color, caption=""):
        ws.row_dimensions[r].height = 22
        ws.row_dimensions[r + 1].height = 42
        ws.row_dimensions[r + 2].height = 20
        merge(ws, r, c1, r, c2, value=label, f=font(8.5, True, "FFFFFF"),
              fl=fill(color), al=Alignment(horizontal="left", vertical="bottom", indent=1))
        merge(ws, r + 1, c1, r + 1, c2, value=formula, f=font(24, True, "FFFFFF"),
              fl=fill(color), al=Alignment(horizontal="left", vertical="center", indent=1),
              nf=nf)
        merge(ws, r + 2, c1, r + 2, c2, value=caption, f=font(8, False, "E8EEFA", italic=True),
              fl=fill(color), al=Alignment(horizontal="left", vertical="top", indent=1))

    def section(self, ws, r, c1, c2, text):
        ws.row_dimensions[r].height = 24
        merge(ws, r, c1, r, c2, value=text, f=font(10.5, True, NAVY),
              al=Alignment(vertical="bottom"), bd=Border(bottom=side(NAVY, "medium")))

    # ===================================================================== #
    # CONTROL PANEL
    # ===================================================================== #
    def build_control(self):
        ws = self.sheet(S_CP)
        last = self.widths(ws, [34, 20, 44, 3, 20, 20, 24, 24])
        self.chrome(ws, S_CP, "Every global setting lives here. Change a yellow cell once "
                    "and all sheets, statuses, orders and dashboards recalculate.",
                    last, "CONTROL PANEL")
        ws.freeze_panes = "A5"

        rows = []

        def sec(title):
            rows.append(("sec", title))

        def setting(key, label, value, nf=None, note="", kind="set", dvs=None):
            rows.append(("set", key, label, value, nf, note, kind, dvs))

        sec("GENERAL")
        setting("cfg_Site", "Site / office name", self.site, None,
                "Shown in every title bar.")
        setting("cfg_AsOf", "Report date (as of)",
                self.as_of if self.as_of else "=TODAY()", F_DATE,
                "Leave as =TODAY() for live reporting, or type a date to report as of that day.")
        sec("SYSTEM 1  ·  INVENTORY RULES")
        setting("cfg_Low", "Low-stock threshold (% of PAR)", 0.5, F_PCT,
                "At or below this share of PAR an item turns LOW (amber).",
                dvs=("decimal", "0", "1"))
        setting("cfg_Crit", "Critical threshold (% of PAR)", 0.25, F_PCT,
                "At or below this share of PAR an item turns CRITICAL (red).",
                dvs=("decimal", "0", "1"))
        setting("cfg_Lead", "Supplier lead time (days)", 7, F_INT,
                "Used to calculate Need-By dates on the order list.",
                dvs=("whole", "0", "120"))
        setting("cfg_ExpiryWarn", "Expiry warning window (days)", 30, F_INT,
                "Items expiring within this many days are flagged EXPIRING.",
                dvs=("whole", "0", "365"))
        setting("cfg_OrderMonth", "Order month",
                self.order_month or "=DATE(YEAR(cfg_AsOf),MONTH(cfg_AsOf),1)", F_MONTH,
                "Defaults to the report month automatically; type a date to plan another month.")
        setting("cfg_ReorderOK", "Also reorder OK items below PAR?", "No", None,
                "No = order only LOW and CRITICAL items. Yes = top every item up to PAR.",
                dvs=("list", '"Yes,No"', None))
        sec("SYSTEM 2  ·  PANTRY PILOT")
        setting("cfg_PilotStart", "Pilot start date", self.pilot_start, "ddd, mmm d, yyyy",
                "Day 1 of the pilot.")
        setting("cfg_PilotDays", "Pilot duration (days)", self.pilot_days, F_INT,
                "Calendar days, including the start and end dates.",
                dvs=("whole", "1", "730"))
        setting("cfg_PilotEnd", "Pilot end date", "=cfg_PilotStart+cfg_PilotDays-1",
                "ddd, mmm d, yyyy", "Calculated: start + duration − 1.", kind="calc")
        setting("cfg_PilotWeeks", "Pilot weeks", "=ROUNDUP(cfg_PilotDays/7,0)", F_INT,
                "Calculated.", kind="calc")
        sec("PILOT DECISION RULES  (drive the leadership recommendation)")
        setting("cfg_MinFeedback", "Minimum feedback entries for a decision", 10, F_INT,
                "Below this the dashboard says GATHER MORE DATA.",
                dvs=("whole", "0", "10000"))
        setting("cfg_ExpandRating", "EXPAND if average rating ≥", 4.2, "0.0",
                "…and positive share at or above the next setting.",
                dvs=("decimal", "1", "5"))
        setting("cfg_ExpandPositive", "EXPAND if positive feedback ≥", 0.6, F_PCT, "",
                dvs=("decimal", "0", "1"))
        setting("cfg_ContinueRating", "CONTINUE if average rating ≥", 3.5, "0.0",
                "Below this the recommendation is ADJUST.", dvs=("decimal", "1", "5"))

        r = FIRST - 1
        for row in rows:
            r += 1
            if row[0] == "sec":
                r += 1 if r > FIRST else 0
                self.section(ws, r, 2, 4, row[1])
                continue
            _, key, label, value, nf, note, kind, dvs = row
            ws.row_dimensions[r].height = 22
            style(ws.cell(r, 2, label), f=font(10, True, "1F2A44"), al=LEFT, bd=BOTTOM)
            v = ws.cell(r, 3, value)
            if kind == "set":
                style(v, f=font(10.5, True, SET_FONT), fl=fill(SET_FILL), al=CENTER, nf=nf,
                      unlock=True, bd=BOX)
            else:
                style(v, f=font(10.5, True, NAVY), fl=fill(PANEL), al=CENTER, nf=nf, bd=BOX)
            style(ws.cell(r, 4, note), f=font(9, False, SLATE, italic=True), al=LEFT, bd=BOTTOM)
            self.name(key, cell_ref(S_CP, f"C{r}"))
            if dvs:
                self.dv(ws, f"C{r}", dvs[0], dvs[1], dvs[2])
        end_settings = r

        # --- lists (right side) ---------------------------------------------- #
        lists = [
            ("StaffList", 6, "STAFF", "Names for Counted By / Refilled By / Owner",
             (self.demo or {}).get("staff", [])),
            ("TeamList", 7, "AREAS / TEAMS", "Used by the Feedback Log",
             (self.demo or {}).get("teams", ["Front Office", "Operations", "Administration"])),
            ("SnackCats", 8, "SNACK CATEGORIES", "Source categories placed in the pantry as Snack",
             self.snack_cats),
            ("BevCats", 9, "BEVERAGE CATEGORIES", "Source categories placed as Beverage",
             self.bev_cats),
        ]
        top = FIRST
        self.section(ws, top, 6, 9, "LISTS  (dropdown sources — edit anytime)")
        for nm, c, head, hint, values in lists:
            style(ws.cell(top + 1, c, head), f=font(8.5, True, "FFFFFF"), fl=fill(SET_HDR),
                  al=CENTER)
            style(ws.cell(top + 2, c, hint), f=font(8, False, SLATE, italic=True), al=CENTER)
            ws.row_dimensions[top + 2].height = 34
            r1, r2 = top + 3, top + 2 + CAP_LIST
            for i in range(CAP_LIST):
                cell = ws.cell(r1 + i, c, values[i] if i < len(values) else None)
                style(cell, f=font(10, False, SET_FONT), fl=fill(SET_FILL), al=LEFT,
                      unlock=True, bd=Border(bottom=side("FFFFFF")))
            L = col_letter(c)
            if nm in ("SnackCats", "BevCats"):
                self.name(nm, rng(S_CP, L, r1, r2))
                self.dv(ws, f"{L}{r1}:{L}{r2}", "list", "=CategoryList",
                        prompt="Pick a category from the Source Data.", strict=False,
                        title="Category")
            else:
                self.name(nm, f"OFFSET('{S_CP}'!${L}${r1},0,0,"
                              f"MAX(1,COUNTA('{S_CP}'!${L}${r1}:${L}${r2})),1)")

        # legend
        lr = max(end_settings, top + 3 + CAP_LIST) + 2
        self.section(ws, lr, 2, 4, "HOW THIS WORKBOOK IS COLOR-CODED")
        legend = [
            (IN_FILL, IN_FONT, "Blue cells", "Your weekly entries: physical counts, refills, feedback."),
            (SET_FILL, SET_FONT, "Yellow cells", "One-time setup: settings, lists and the source product file."),
            (PANEL, NAVY, "Grey / white cells", "Calculated automatically. Locked — never type here."),
        ]
        for i, (bg, fg, a, b) in enumerate(legend):
            rr = lr + 1 + i
            ws.row_dimensions[rr].height = 22
            style(ws.cell(rr, 2, a), f=font(10, True, fg), fl=fill(bg), al=LEFT, bd=BOX)
            merge(ws, rr, 3, rr, 4, value=b, f=font(9.5, False, SLATE), al=LEFT)
        self.protect(ws)

    # ===================================================================== #
    # SOURCE DATA (swap zone)
    # ===================================================================== #
    def build_source(self):
        ws = self.sheet(S_SRC)
        last = self.widths(ws, [24, 44, 14, 14, 14, 10, 20])
        self.chrome(ws, S_SRC, "Product master from the source file. To launch another office: "
                    "clear this table and paste Category, Item Name, Item #, Recent Date Received, "
                    "Expiration Date and PAR Value from the new source file (values only, no blank rows).",
                    last, "SOURCE DATA")
        heads = ["Category", "Item Name", "Item #", "Recent Date Received", "Expiration Date",
                 "PAR Value", "Pilot Type Override"]
        self.headers(ws, heads, {h: "set" for h in heads})
        r2 = HDR + CAP_PRODUCTS
        for i, col in enumerate(range(2, 9)):
            nf = {4: F_DATE, 5: F_DATE, 6: F_INT}.get(col - 1)
            self.body(ws, col, FIRST, r2, None, nf=nf, kind="set",
                      al=LEFT if col in (2, 3) else CENTER)
        for i, p in enumerate(self.products[:CAP_PRODUCTS]):
            r = FIRST + i
            ws.cell(r, 2, p["category"])
            ws.cell(r, 3, p["name"])
            ws.cell(r, 4, p["number"])
            ws.cell(r, 5, p["received"])
            ws.cell(r, 6, p.get("expiration"))
            ws.cell(r, 7, p["par"])
        self.dv(ws, f"H{FIRST}:H{r2}", "list", '"Snack,Beverage,Not in Pilot"',
                prompt="Optional. Overrides the category mapping on the Control Panel.",
                title="Pilot type")
        self.dv(ws, f"G{FIRST}:G{r2}", "decimal", "0", "1000000")
        self.table(ws, "tblSource", 2, 8, r2)
        ws.freeze_panes = f"A{FIRST}"
        self.name("SrcCount", f"COUNTA('{S_SRC}'!$C${FIRST}:$C${r2})")

    # ===================================================================== #
    # ENGINE (hidden helper calculations)
    # ===================================================================== #
    def build_engine(self):
        ws = self.sheet(S_EN)
        ws.sheet_state = "hidden"
        n = CAP_PRODUCTS
        e1, e2 = 2, 1 + n
        heads = {1: "Idx", 2: "Category", 3: "Item Name", 4: "Item #", 5: "Label",
                 6: "Received", 7: "PAR", 8: "CatIdx", 9: "PilotType", 10: "PilotIdx",
                 11: "WatchScore", 12: "Expiry", 28: "ExpiryScore",
                 13: "k", 14: "Category", 15: "OrderUnits", 16: "Score", 17: "ItemsToOrder",
                 19: "k", 20: "SrcRow", 21: "PilotLabel", 22: "PilotType", 23: "Units",
                 24: "Score", 26: "FeedbackProducts"}
        for c, h in heads.items():
            ws.cell(1, c, h).font = font(9, True)
        S = f"'{S_SRC}'!"
        M = f"'{S_MS}'!"
        for i in range(n):
            r = e1 + i
            s = FIRST + i
            m = FIRST + i
            ws.cell(r, 1, i + 1)
            ws.cell(r, 2, f'=IF({S}$C{s}="","",IF({S}B{s}="","Uncategorized",{S}B{s}))')
            ws.cell(r, 3, f'=IF({S}$C{s}="","",TRIM({S}C{s}))')
            ws.cell(r, 4, f'=IF(OR({S}$C{s}="",{S}D{s}=""),"",{S}D{s})')
            ws.cell(r, 5, f'=IF(C{r}="","",IF(COUNTIF($C${e1}:$C${e2},C{r})>1,'
                          f'C{r}&" (#"&D{r}&")",C{r}))')
            ws.cell(r, 6, f'=IF(OR({S}$C{s}="",{S}E{s}=""),"",{S}E{s})')
            ws.cell(r, 7, f'=IF(OR({S}$C{s}="",{S}G{s}=""),"",{S}G{s})')
            ws.cell(r, 8, f'=IF(B{r}="","",IF(COUNTIF($B${e1}:B{r},B{r})=1,MAX($H$1:H{r - 1})+1,""))')
            ws.cell(r, 9, f'=IF(C{r}="","",IF({S}H{s}="Not in Pilot","",IF({S}H{s}<>"",{S}H{s},'
                          f'IF(COUNTIF(SnackCats,B{r})>0,"Snack",IF(COUNTIF(BevCats,B{r})>0,"Beverage","")))))')
            ws.cell(r, 10, f'=IF(I{r}="","",COUNTIF($I${e1}:I{r},"Snack")+COUNTIF($I${e1}:I{r},"Beverage"))')
            ws.cell(r, 11, f'=IF(OR({M}J{m}="CRITICAL",{M}J{m}="LOW"),{M}I{m}+A{r}/1000000,"")')
            ws.cell(r, 12, f'=IF(OR({S}$C{s}="",{S}F{s}=""),"",{S}F{s})')
            ws.cell(r, 28, f'=IF(ISNUMBER(L{r}),L{r}+A{r}/1000000,"")')
        self.name("Eng_Cat", rng(S_EN, "B", e1, e2))
        self.name("Eng_Name", rng(S_EN, "C", e1, e2))
        self.name("Eng_Item", rng(S_EN, "D", e1, e2))
        self.name("Eng_Label", rng(S_EN, "E", e1, e2))
        self.name("Eng_PAR", rng(S_EN, "G", e1, e2))
        self.name("Eng_PilotIdx", rng(S_EN, "J", e1, e2))
        self.name("Eng_Watch", rng(S_EN, "K", e1, e2))
        self.name("Eng_Expiry", rng(S_EN, "AB", e1, e2))
        self.name("ProductList", f"OFFSET('{S_EN}'!$E$2,0,0,MAX(1,SrcCount),1)")

        # categories
        c2 = 1 + CAP_CATS
        for k in range(1, CAP_CATS + 1):
            r = 1 + k
            ws.cell(r, 13, k)
            ws.cell(r, 14, f'=IFERROR(INDEX(Eng_Cat,MATCH(M{r},$H${e1}:$H${e2},0)),"")')
            ws.cell(r, 15, f'=IF(N{r}="","",SUMIFS(OL_Qty,OL_Cat,N{r}))')
            ws.cell(r, 16, f'=IF(N{r}="",-1,O{r}+(1000-M{r})/10000)')
            ws.cell(r, 17, f'=IF(N{r}="","",COUNTIFS(OL_Cat,N{r},OL_Flag,"ORDER"))')
        self.name("Cat_Name", rng(S_EN, "N", 2, c2))
        self.name("Cat_Units", rng(S_EN, "O", 2, c2))
        self.name("Cat_Score", rng(S_EN, "P", 2, c2))
        self.name("CategoryList", f"OFFSET('{S_EN}'!$N$2,0,0,MAX(1,MAX('{S_EN}'!$H${e1}:$H${e2})),1)")

        # pilot products
        p2 = 1 + CAP_PILOT
        for k in range(1, CAP_PILOT + 1):
            r = 1 + k
            ws.cell(r, 19, k)
            ws.cell(r, 20, f'=IFERROR(MATCH(S{r},Eng_PilotIdx,0),"")')
            ws.cell(r, 21, f'=IF(T{r}="","",INDEX(Eng_Label,T{r}))')
            ws.cell(r, 22, f'=IF(T{r}="","",INDEX(\'{S_EN}\'!$I${e1}:$I${e2},T{r}))')
            ws.cell(r, 23, f'=IF(U{r}="","",SUMIFS(RL_Added,RL_Prod,U{r}))')
            ws.cell(r, 24, f'=IF(U{r}="",-1,W{r}+(1000-S{r})/10000)')
        self.name("Pil_Row", rng(S_EN, "T", 2, p2))
        self.name("Pil_Label", rng(S_EN, "U", 2, p2))
        self.name("Pil_Type", rng(S_EN, "V", 2, p2))
        self.name("Pil_Units", rng(S_EN, "W", 2, p2))
        self.name("Pil_Score", rng(S_EN, "X", 2, p2))
        self.name("PilotCount", f"MAX(Eng_PilotIdx)")
        self.name("PilotList", f"OFFSET('{S_EN}'!$U$2,0,0,MAX(1,PilotCount),1)")

        ws.cell(2, 26, "Pantry — General")
        for k in range(1, CAP_PILOT + 1):
            ws.cell(2 + k, 26, f'=U{1 + k}')
        self.name("FeedbackProducts", f"OFFSET('{S_EN}'!$Z$2,0,0,1+PilotCount,1)")
        self.protect(ws)

    # ===================================================================== #
    # INVENTORY MASTER
    # ===================================================================== #
    def build_master(self):
        ws = self.sheet(S_MS)
        heads = ["Category", "Item Name", "Item #", "PAR Level", "Starting Qty", "Latest Count",
                 "Current Stock", "Stock %", "Status", "Suggested Order Qty", "Last Count Date",
                 "Last Received Date", "Counts Logged", "Expiration Date", "Days to Expiry",
                 "Expiry Status"]
        last = self.widths(ws, [20, 34, 11, 10, 11, 11, 11, 14, 12, 13, 15, 15, 10, 14, 10, 12])
        self.chrome(ws, S_MS, "One row per product, built from Source Data. Latest counts are read "
                    "automatically from the Weekly Count log. Status: OK › LOW › CRITICAL using the "
                    "Control Panel thresholds. Nothing on this page is typed by hand.",
                    last, "INVENTORY MANAGEMENT")
        self.headers(ws, heads)
        r2 = HDR + CAP_PRODUCTS
        E = f"'{S_EN}'!"
        cols = [
            ("B", f"={E}B{{e}}", None, LEFT),
            ("C", f"={E}C{{e}}", None, LEFT),
            ("D", f"={E}D{{e}}", None, None),
            ("E", f"={E}G{{e}}", F_INT, None),
            ("F", '=IF(OR($C{r}="",N{r}=0),"",INDEX(CL_Count,MATCH(' + E + 'E{e}&"|1",CL_Key,0)))', F_INT, None),
            ("G", '=IF(OR($C{r}="",N{r}=0),"",INDEX(CL_Count,MATCH(' + E + 'E{e}&"|"&N{r},CL_Key,0)))', F_INT, None),
            ("H", '=IF($C{r}="","",G{r})', F_INT, None),
            ("I", '=IF(OR(H{r}="",E{r}="",N(E{r})=0),"",H{r}/E{r})', F_PCT, None),
            ("J", '=IF($C{r}="","",IF(H{r}="","NO COUNT",IF(OR(E{r}="",N(E{r})=0),"NO PAR",'
                  'IF(I{r}<=cfg_Crit,"CRITICAL",IF(I{r}<=cfg_Low,"LOW","OK")))))', None, None),
            ("K", '=IF(OR($C{r}="",H{r}="",E{r}=""),"",MAX(0,E{r}-H{r}))', F_INT, None),
            ("L", '=IF(OR($C{r}="",N{r}=0),"",INDEX(CL_Date,MATCH(' + E + 'E{e}&"|"&N{r},CL_Key,0)))', F_DATE, None),
            ("M", f"={E}F{{e}}", F_DATE, None),
            ("N", '=IF($C{r}="","",COUNTIFS(CL_Prod,' + E + 'E{e},CL_Count,"<>"))', F_INT, None),
            ("O", f"={E}L{{e}}", F_DATE, None),
            ("P", '=IF(ISNUMBER(O{r}),O{r}-cfg_AsOf,"")', "0;[Red]-0;0", None),
            ("Q", '=IF($C{r}="","",IF(NOT(ISNUMBER(O{r})),"N/A",IF(P{r}<0,"EXPIRED",'
                  'IF(P{r}<=cfg_ExpiryWarn,"EXPIRING","OK"))))', None, None),
        ]
        for letter, f, nf, al in cols:
            c = ord(letter) - 64
            for r in range(FIRST, r2 + 1):
                cell = ws.cell(r, c, f.format(r=r, e=r - FIRST + 2))
                style(cell, f=font(10), al=al or CENTER, nf=nf)
        self.band(ws, 3, 2, 17, FIRST, r2)
        self.status_colors(ws, f"Q{FIRST}:Q{r2}", {
            "EXPIRED": (CRIT_BG, CRIT_FG), "EXPIRING": (LOW_BG, LOW_FG), "OK": (OK_BG, OK_FG),
            "N/A": (NONE_BG, NONE_FG)})
        self.status_colors(ws, f"J{FIRST}:J{r2}", {
            "OK": (OK_BG, OK_FG), "LOW": (LOW_BG, LOW_FG), "CRITICAL": (CRIT_BG, CRIT_FG),
            "NO COUNT": (NONE_BG, NONE_FG), "NO PAR": (NONE_BG, NONE_FG)})
        ws.conditional_formatting.add(f"I{FIRST}:I{r2}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color="8FA9EA"))
        self.table(ws, "tblInventory", 2, 17, r2)
        ws.freeze_panes = f"D{FIRST}"
        for nm, L in [("MS_Name", "C"), ("MS_Status", "J"), ("MS_Stock", "H"), ("MS_PAR", "E"),
                      ("MS_Pct", "I"), ("MS_Sugg", "K"), ("MS_Cat", "B"), ("MS_Exp", "O"),
                      ("MS_ExpDays", "P"), ("MS_ExpStatus", "Q")]:
            self.name(nm, rng(S_MS, L, FIRST, r2))
        self.protect(ws)

    # ===================================================================== #
    # WEEKLY COUNT LOG
    # ===================================================================== #
    def build_count(self):
        ws = self.sheet(S_CL)
        heads = ["Count Date", "Week #", "Product", "Item #", "Physical Count", "Prior Count",
                 "Weekly Usage", "Counted By", "PAR Level", "Variance to PAR", "Status", "Month",
                 "Record ID", "Key", "Seq"]
        last = self.widths(ws, [14, 8, 36, 11, 12, 11, 11, 16, 10, 12, 12, 11, 12, 6, 6], hidden=("O", "P"))
        self.chrome(ws, S_CL, "Weekly physical count — type only the BLUE columns: date, pick the "
                    "product, the quantity on the shelf and who counted. Add new counts at the bottom "
                    "(oldest first). Everything else fills itself and feeds the Inventory Master.",
                    last, "INVENTORY MANAGEMENT")
        kinds = {h: "in" for h in ["Count Date", "Product", "Physical Count", "Counted By"]}
        self.headers(ws, heads, kinds)
        r2 = HDR + CAP_COUNTS
        B = self.body
        B(ws, 2, FIRST, r2, None, F_DATE, kind="in")
        B(ws, 3, FIRST, r2, '=IF(B{r}="","",WEEKNUM(B{r},21))', "0")
        B(ws, 4, FIRST, r2, None, None, kind="in", al=LEFT)
        B(ws, 5, FIRST, r2, '=IF(D{r}="","",IFERROR(INDEX(Eng_Item,MATCH(D{r},Eng_Label,0)),"not in list"))')
        B(ws, 6, FIRST, r2, None, F_INT, kind="in")
        B(ws, 7, FIRST, r2, '=IF(OR(P{r}="",N(P{r})<=1),"",INDEX($F$' + str(FIRST) + ':$F$' + str(r2)
          + ',MATCH(D{r}&"|"&(P{r}-1),$O$' + str(FIRST) + ':$O$' + str(r2) + ',0)))', F_INT)
        B(ws, 8, FIRST, r2, '=IF(G{r}="","",MAX(0,G{r}-F{r}))', F_INT)
        B(ws, 9, FIRST, r2, None, None, kind="in", al=LEFT)
        B(ws, 10, FIRST, r2, '=IF(D{r}="","",IFERROR(INDEX(Eng_PAR,MATCH(D{r},Eng_Label,0)),""))', F_INT)
        B(ws, 11, FIRST, r2, '=IF(OR(F{r}="",J{r}=""),"",F{r}-J{r})', '+#,##0;-#,##0;0')
        B(ws, 12, FIRST, r2, '=IF(OR(F{r}="",J{r}="",N(J{r})=0),"",IF(F{r}/J{r}<=cfg_Crit,"CRITICAL",'
          'IF(F{r}/J{r}<=cfg_Low,"LOW","OK")))')
        B(ws, 13, FIRST, r2, '=IF(B{r}="","",DATE(YEAR(B{r}),MONTH(B{r}),1))', F_MONTH)
        B(ws, 14, FIRST, r2, '=IF(OR(B{r}="",D{r}=""),"","CNT-"&TEXT(ROW()-' + str(HDR) + ',"00000"))')
        B(ws, 15, FIRST, r2, '=IF(P{r}="","",D{r}&"|"&P{r})')
        B(ws, 16, FIRST, r2, '=IF(OR(D{r}="",F{r}=""),"",COUNTIFS($D$' + str(FIRST) + ':D{r},D{r},$F$'
          + str(FIRST) + ':F{r},"<>"))')
        for cols in ["C", "E", "G:H", "J:N"]:
            a, b = (cols.split(":") + [cols])[:2]
            self.band(ws, 4, ord(a) - 64, ord(b) - 64, FIRST, r2)
        self.status_colors(ws, f"L{FIRST}:L{r2}", {
            "OK": (OK_BG, OK_FG), "LOW": (LOW_BG, LOW_FG), "CRITICAL": (CRIT_BG, CRIT_FG)})
        ws.conditional_formatting.add(f"E{FIRST}:E{r2}", CellIsRule(
            operator="equal", formula=['"not in list"'], font=Font(color=C_RED, bold=True)))

        self.dv(ws, f"B{FIRST}:B{r2}", "date", "DATE(2020,1,1)", "DATE(2100,12,31)",
                prompt="Date of the physical count.", title="Count date")
        self.dv(ws, f"D{FIRST}:D{r2}", "list", "=ProductList",
                prompt="Pick the product from the list.", title="Product")
        self.dv(ws, f"F{FIRST}:F{r2}", "decimal", "0", "1000000",
                prompt="Units physically on the shelf.", title="Physical count")
        self.dv(ws, f"I{FIRST}:I{r2}", "list", "=StaffList", strict=False)
        self.table(ws, "tblCounts", 2, 16, r2)
        ws.freeze_panes = f"A{FIRST}"
        for nm, L in [("CL_Date", "B"), ("CL_Prod", "D"), ("CL_Count", "F"), ("CL_Key", "O")]:
            self.name(nm, rng(S_CL, L, FIRST, r2))
        if self.demo:
            for i, (d, prod, qty, who) in enumerate(self.demo["counts"]):
                r = FIRST + i
                ws.cell(r, 2, d); ws.cell(r, 4, prod); ws.cell(r, 6, qty); ws.cell(r, 9, who)
        self.protect(ws)

    # ===================================================================== #
    # MONTHLY ORDER LIST
    # ===================================================================== #
    def build_orders(self):
        ws = self.sheet(S_OL)
        heads = ["Category", "Item Name", "Item #", "Current Stock", "PAR Level", "Order Qty",
                 "Status", "Order Month", "Need-By Date", "Priority", "Order Flag"]
        last = self.widths(ws, [20, 34, 11, 12, 10, 11, 12, 12, 14, 12, 13])
        self.chrome(ws, S_OL, "Auto-generated purchase list. Use the header filters (e.g. Order Flag = "
                    "ORDER, Priority = URGENT). Order Qty = PAR − Current Stock, never negative. "
                    "Nothing on this page is typed by hand.", last, "INVENTORY MANAGEMENT")
        self.headers(ws, heads)
        r2 = HDR + CAP_PRODUCTS
        M = f"'{S_MS}'!"
        B = self.body
        B(ws, 2, FIRST, r2, f"={M}B{{r}}", al=LEFT)
        B(ws, 3, FIRST, r2, f"={M}C{{r}}", al=LEFT)
        B(ws, 4, FIRST, r2, f"={M}D{{r}}")
        B(ws, 5, FIRST, r2, f"={M}H{{r}}", F_INT)
        B(ws, 6, FIRST, r2, f"={M}E{{r}}", F_INT)
        B(ws, 7, FIRST, r2, '=IF($C{r}="","",IF(L{r}="ORDER",' + M + 'K{r},0))', F_INT)
        B(ws, 8, FIRST, r2, f"={M}J{{r}}")
        B(ws, 9, FIRST, r2, '=IF($C{r}="","",cfg_OrderMonth)', F_MONTH)
        B(ws, 10, FIRST, r2, '=IF(OR($C{r}="",L{r}<>"ORDER"),"",IF(K{r}="URGENT",cfg_AsOf+cfg_Lead,'
          'MAX(cfg_OrderMonth,cfg_AsOf)+cfg_Lead))', F_DATE)
        B(ws, 11, FIRST, r2, '=IF($C{r}="","",IF(H{r}="CRITICAL","URGENT",IF(H{r}="LOW","STANDARD",'
          'IF(AND(cfg_ReorderOK="Yes",H{r}="OK",N(' + M + 'K{r})>0),"STANDARD","NONE"))))')
        B(ws, 12, FIRST, r2, '=IF($C{r}="","",IF(AND(K{r}<>"NONE",N(' + M + 'K{r})>0),"ORDER","NO ORDER"))')
        self.band(ws, 3, 2, 12, FIRST, r2)
        self.status_colors(ws, f"H{FIRST}:H{r2}", {
            "OK": (OK_BG, OK_FG), "LOW": (LOW_BG, LOW_FG), "CRITICAL": (CRIT_BG, CRIT_FG),
            "NO COUNT": (NONE_BG, NONE_FG), "NO PAR": (NONE_BG, NONE_FG)})
        self.status_colors(ws, f"K{FIRST}:K{r2}", {
            "URGENT": (CRIT_BG, CRIT_FG), "STANDARD": (LOW_BG, LOW_FG), "NONE": (NONE_BG, NONE_FG)})
        self.status_colors(ws, f"L{FIRST}:L{r2}", {
            "ORDER": (NAVY, "FFFFFF"), "NO ORDER": (PAGE, MUTED)})
        ws.conditional_formatting.add(f"G{FIRST}:G{r2}", CellIsRule(
            operator="greaterThan", formula=["0"], font=Font(bold=True, color=NAVY)))
        self.table(ws, "tblOrders", 2, 12, r2)
        ws.freeze_panes = f"D{FIRST}"
        for nm, L in [("OL_Cat", "B"), ("OL_Qty", "G"), ("OL_Flag", "L"), ("OL_Priority", "K")]:
            self.name(nm, rng(S_OL, L, FIRST, r2))
        self.protect(ws)

    # ===================================================================== #
    # INVENTORY DASHBOARD
    # ===================================================================== #
    def build_inv_dash(self):
        ws = self.sheet(S_ID)
        last = self.widths(ws, [10.5] * 16)
        self.chrome(ws, S_ID, "Live inventory health. Updates the moment a count is logged or a "
                    "threshold changes on the Control Panel.", last, "INVENTORY DASHBOARD")
        ws.freeze_panes = "A4"
        counted = '(COUNTIF(MS_Status,"OK")+COUNTIF(MS_Status,"LOW")+COUNTIF(MS_Status,"CRITICAL"))'
        cards = [
            ("TOTAL ITEMS", '=SUMPRODUCT(--(MS_Name<>""))', F_INT, C_NAVY,
             '=COUNTIF(MS_Status,"NO COUNT")&" not counted yet"'),
            ("OK ITEMS", '=COUNTIF(MS_Status,"OK")', F_INT, C_GREEN,
             '="Above "&FIXED(cfg_Low*100,0)&"% of PAR"'),
            ("LOW ITEMS", '=COUNTIF(MS_Status,"LOW")', F_INT, C_AMBER,
             '="≤ "&FIXED(cfg_Low*100,0)&"% of PAR"'),
            ("CRITICAL ITEMS", '=COUNTIF(MS_Status,"CRITICAL")', F_INT, C_RED,
             '="≤ "&FIXED(cfg_Crit*100,0)&"% of PAR"'),
            ("UNITS TO ORDER", "=SUM(OL_Qty)", F_INT, C_BLUE, '="For the order month"'),
            ("ITEMS TO ORDER", '=COUNTIF(OL_Flag,"ORDER")', F_INT, C_TEAL,
             '=COUNTIF(OL_Priority,"URGENT")&" urgent"'),
            ("LAST COUNT DATE", '=IF(MAX(CL_Date)=0,"No counts",MAX(CL_Date))', F_DATE, C_PURPLE,
             '=IF(MAX(CL_Date)=0,"Log a count to start",(cfg_AsOf-MAX(CL_Date))&" days ago")'),
            ("STOCK HEALTH", f'=IF({counted}=0,"—",COUNTIF(MS_Status,"OK")/{counted})', F_PCT, C_NAVY,
             '="Share of counted items at OK"'),
        ]
        for i, (lab, f, nf, colr, cap) in enumerate(cards):
            r = 5 if i < 4 else 9
            c = 2 + (i % 4) * 4
            self.card(ws, r, c, c + 3, lab, f, nf, colr, cap)
        ws.row_dimensions[8].height = 10

        # chart data blocks
        self.section(ws, 13, 2, 7, "STATUS MIX")
        self.section(ws, 13, 9, 17, f"ORDER UNITS BY CATEGORY  (top {TOP_N})")
        tr = 32  # data tables start below the charts
        self.section(ws, tr, 2, 7, "STATUS MIX — DATA")
        style(ws.cell(tr + 1, 2, "Status"), f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        merge(ws, tr + 1, 3, tr + 1, 4, "Items", f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        merge(ws, tr + 1, 5, tr + 1, 7, "Share", f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        mix = [("OK", C_GREEN), ("LOW", C_AMBER), ("CRITICAL", C_RED), ("NO COUNT", "B8C0D2")]
        for i, (s, colr) in enumerate(mix):
            r = tr + 2 + i
            style(ws.cell(r, 2, s), f=font(10, True, colr if s != "NO COUNT" else SLATE), al=CENTER, bd=BOTTOM)
            merge(ws, r, 3, r, 4, f'=COUNTIF(MS_Status,B{r})', f=font(10), al=CENTER, bd=BOTTOM, nf=F_INT)
            merge(ws, r, 5, r, 7, f'=IF(SUM($C${tr + 2}:$C${tr + 5})=0,"",C{r}/SUM($C${tr + 2}:$C${tr + 5}))',
                  f=font(10), al=CENTER, bd=BOTTOM, nf=F_PCT)

        self.section(ws, tr, 9, 17, "ORDER UNITS BY CATEGORY — DATA")
        style(ws.cell(tr + 1, 9, "#"), f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        merge(ws, tr + 1, 10, tr + 1, 13, "Category", f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        merge(ws, tr + 1, 14, tr + 1, 15, "Order Units", f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        merge(ws, tr + 1, 16, tr + 1, 17, "Items to Order", f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        for k in range(1, TOP_N + 1):
            r = tr + 1 + k
            style(ws.cell(r, 9, k), f=font(9, False, MUTED), al=CENTER, bd=BOTTOM)
            idx = f'MATCH(LARGE(Cat_Score,I{r}),Cat_Score,0)'
            merge(ws, r, 10, r, 13, f'=IF(LARGE(Cat_Score,I{r})<0,"",INDEX(Cat_Name,{idx}))',
                  f=font(10), al=LEFT, bd=BOTTOM)
            merge(ws, r, 14, r, 15, f'=IF(J{r}="","",INDEX(Cat_Units,{idx}))', f=font(10, True, NAVY),
                  al=CENTER, bd=BOTTOM, nf=F_INT)
            merge(ws, r, 16, r, 17, f'=IF(J{r}="","",INDEX(\'{S_EN}\'!$Q$2:$Q${1 + CAP_CATS},{idx}))',
                  f=font(10), al=CENTER, bd=BOTTOM, nf=F_INT)

        # watchlist
        wr = tr + TOP_N + 4
        self.section(ws, wr, 2, 17, "CRITICAL & LOW WATCHLIST  (lowest stock % first)")
        wh = [(2, 2, "#"), (3, 6, "Item"), (7, 9, "Category"), (10, 11, "Stock %"),
              (12, 13, "On Hand"), (14, 15, "Order Qty"), (16, 17, "Status")]
        for c1, c2, h in wh:
            merge(ws, wr + 1, c1, wr + 1, c2, h, f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        for k in range(1, TOP_N + 1):
            r = wr + 1 + k
            style(ws.cell(r, 2, k), f=font(9, False, MUTED), al=CENTER, bd=BOTTOM)
            idx = f'MATCH(SMALL(Eng_Watch,B{r}),Eng_Watch,0)'
            has = f'COUNT(Eng_Watch)>=B{r}'
            merge(ws, r, 3, r, 6, f'=IF({has},INDEX(Eng_Label,{idx}),"")', f=font(10), al=LEFT, bd=BOTTOM)
            merge(ws, r, 7, r, 9, f'=IF({has},INDEX(MS_Cat,{idx}),"")', f=font(10, False, SLATE), al=LEFT, bd=BOTTOM)
            merge(ws, r, 10, r, 11, f'=IF({has},INDEX(MS_Pct,{idx}),"")', f=font(10, True), al=CENTER, bd=BOTTOM, nf=F_PCT)
            merge(ws, r, 12, r, 13, f'=IF({has},INDEX(MS_Stock,{idx}),"")', f=font(10), al=CENTER, bd=BOTTOM, nf=F_INT)
            merge(ws, r, 14, r, 15, f'=IF({has},INDEX(MS_Sugg,{idx}),"")', f=font(10, True, NAVY), al=CENTER, bd=BOTTOM, nf=F_INT)
            merge(ws, r, 16, r, 17, f'=IF({has},INDEX(MS_Status,{idx}),"")', f=font(10, True), al=CENTER, bd=BOTTOM)
        area = f"P{wr + 2}:Q{wr + 1 + TOP_N}"
        self.status_colors(ws, area, {"LOW": (LOW_BG, LOW_FG), "CRITICAL": (CRIT_BG, CRIT_FG)})
        ws.cell(wr + TOP_N + 2, 2, "No item is LOW or CRITICAL when this list is empty.").font = \
            font(8.5, False, MUTED, italic=True)

        # expiration watch
        er = wr + TOP_N + 4
        self.section(ws, er, 2, 17, '="EXPIRATION WATCH  (soonest first)   ·   "&COUNTIF(MS_ExpStatus,"EXPIRED")'
                     '&" expired   ·   "&COUNTIF(MS_ExpStatus,"EXPIRING")&" expiring within "&cfg_ExpiryWarn&" days"')
        eh = [(2, 2, "#"), (3, 6, "Item"), (7, 9, "Category"), (10, 11, "Expiration"),
              (12, 13, "Days Left"), (14, 15, "On Hand"), (16, 17, "Expiry Status")]
        for c1, c2, h in eh:
            merge(ws, er + 1, c1, er + 1, c2, h, f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        for k in range(1, TOP_N + 1):
            r = er + 1 + k
            style(ws.cell(r, 2, k), f=font(9, False, MUTED), al=CENTER, bd=BOTTOM)
            idx = f'MATCH(SMALL(Eng_Expiry,B{r}),Eng_Expiry,0)'
            has = f'COUNT(Eng_Expiry)>=B{r}'
            merge(ws, r, 3, r, 6, f'=IF({has},INDEX(Eng_Label,{idx}),"")', f=font(10), al=LEFT, bd=BOTTOM)
            merge(ws, r, 7, r, 9, f'=IF({has},INDEX(MS_Cat,{idx}),"")', f=font(10, False, SLATE), al=LEFT, bd=BOTTOM)
            merge(ws, r, 10, r, 11, f'=IF({has},INDEX(MS_Exp,{idx}),"")', f=font(10), al=CENTER, bd=BOTTOM, nf=F_DATE)
            merge(ws, r, 12, r, 13, f'=IF({has},INDEX(MS_ExpDays,{idx}),"")', f=font(10, True), al=CENTER, bd=BOTTOM,
                  nf="0;[Red]-0;0")
            merge(ws, r, 14, r, 15, f'=IF({has},INDEX(MS_Stock,{idx}),"")', f=font(10), al=CENTER, bd=BOTTOM, nf=F_INT)
            merge(ws, r, 16, r, 17, f'=IF({has},INDEX(MS_ExpStatus,{idx}),"")', f=font(10, True), al=CENTER, bd=BOTTOM)
        self.status_colors(ws, f"P{er + 2}:Q{er + 1 + TOP_N}", {
            "EXPIRED": (CRIT_BG, CRIT_FG), "EXPIRING": (LOW_BG, LOW_FG), "OK": (OK_BG, OK_FG)})
        ws.cell(er + TOP_N + 2, 2, "Negative days = already expired. Items marked N/A or "
                "Not Available in the source have no date and are not listed.").font = \
            font(8.5, False, MUTED, italic=True)

        # charts
        d = DoughnutChart(holeSize=58)
        d.add_data(Reference(ws, min_col=3, min_row=tr + 2, max_row=tr + 5), titles_from_data=False)
        d.set_categories(Reference(ws, min_col=2, min_row=tr + 2, max_row=tr + 5))
        s = d.series[0]
        for i, (_, colr) in enumerate(mix):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = colr
            pt.graphicalProperties.line.solidFill = "FFFFFF"
            s.dPt.append(pt)
        s.dLbls = DataLabelList(); s.dLbls.showVal = True
        for a in ("showSerName", "showCatName", "showLegendKey", "showPercent", "showLeaderLines"):
            setattr(s.dLbls, a, False)
        d.legend.position = "r"
        d.width, d.height = 11.5, 8.6
        ws.add_chart(d, "B14")

        b = self.hbar(ws, ref_vals=(14, tr + 2, tr + 1 + TOP_N), ref_cats=(10, tr + 2, tr + 1 + TOP_N),
                      color=C_BLUE, title=None)
        b.width, b.height = 17.5, 8.6
        ws.add_chart(b, "I14")
        self.protect(ws)

    def hbar(self, ws, ref_vals, ref_cats, color, title=None):
        b = BarChart()
        b.type = "bar"
        b.style = 10
        b.add_data(Reference(ws, min_col=ref_vals[0], min_row=ref_vals[1], max_row=ref_vals[2]),
                   titles_from_data=False)
        b.set_categories(Reference(ws, min_col=ref_cats[0], min_row=ref_cats[1], max_row=ref_cats[2]))
        b.legend = None
        b.gapWidth = 45
        s = b.series[0]
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
        s.dLbls = DataLabelList(); s.dLbls.showVal = True
        for a in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(s.dLbls, a, False)
        b.x_axis.scaling.orientation = "maxMin"     # rank 1 at the top
        b.x_axis.delete = False
        b.y_axis.delete = True
        b.y_axis.majorGridlines = None
        b.x_axis.majorTickMark = "none"
        b.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE))
        if title:
            b.title = title
        return b

    # ===================================================================== #
    # PILOT PRODUCTS
    # ===================================================================== #
    def build_pilot_products(self):
        ws = self.sheet(S_PP)
        heads = ["Product", "Item #", "Category", "Type", "Pilot Start Qty", "Target Refill Qty",
                 "Active Flag", "Pilot Start", "Pilot End", "First Placed", "Days Active",
                 "Refill Events", "Total Units Placed", "Avg Qty per Refill",
                 "Avg Days Between Refills", "Last Refill", "Units per Active Day"]
        last = self.widths(ws, [32, 10, 18, 11, 11, 11, 13, 13, 13, 13, 9, 9, 11, 11, 12, 16, 11])
        self.chrome(ws, S_PP, "Pantry products — every Source item whose category is mapped to Snack or "
                    "Beverage on the Control Panel (or overridden on Source Data). Start qty = first "
                    "placement; target refill = PAR. All columns are calculated.",
                    last, "PANTRY PILOT")
        self.headers(ws, heads, height=40)
        r2 = HDR + CAP_PILOT
        E = f"'{S_EN}'!"
        B = self.body
        rows = range(FIRST, r2 + 1)
        # write with engine offset
        for r in rows:
            e = r - FIRST + 2
            ws.cell(r, 2).value = f"={E}U{e}"
            ws.cell(r, 3).value = f'=IF($B{r}="","",INDEX(Eng_Item,{E}T{e}))'
            ws.cell(r, 4).value = f'=IF($B{r}="","",INDEX(Eng_Cat,{E}T{e}))'
            ws.cell(r, 5).value = f"={E}V{e}"
            ws.cell(r, 7).value = f'=IF($B{r}="","",INDEX(Eng_PAR,{E}T{e}))'
        B(ws, 2, FIRST, r2, None, al=LEFT)
        B(ws, 3, FIRST, r2, None)
        B(ws, 4, FIRST, r2, None, al=LEFT)
        B(ws, 5, FIRST, r2, None)
        B(ws, 6, FIRST, r2, '=IF(OR($B{r}="",M{r}=0),"",INDEX(RL_After,MATCH($B{r}&"|1",RL_Key,0)))', F_INT)
        B(ws, 7, FIRST, r2, None, F_INT)
        B(ws, 8, FIRST, r2, '=IF($B{r}="","",IF(M{r}=0,"NOT PLACED",IF(cfg_AsOf>cfg_PilotEnd,"ENDED","ACTIVE")))')
        B(ws, 9, FIRST, r2, '=IF($B{r}="","",cfg_PilotStart)', F_DATE)
        B(ws, 10, FIRST, r2, '=IF($B{r}="","",cfg_PilotEnd)', F_DATE)
        B(ws, 11, FIRST, r2, '=IF(OR($B{r}="",M{r}=0),"",INDEX(RL_Date,MATCH($B{r}&"|1",RL_Key,0)))', F_DATE)
        B(ws, 12, FIRST, r2, '=IF(K{r}="","",MAX(0,MIN(cfg_AsOf,cfg_PilotEnd)-K{r}+1))', F_INT)
        B(ws, 13, FIRST, r2, '=IF($B{r}="","",COUNTIFS(RL_Prod,$B{r},RL_Added,"<>"))', F_INT)
        B(ws, 14, FIRST, r2, '=IF($B{r}="","",SUMIFS(RL_Added,RL_Prod,$B{r}))', F_INT)
        B(ws, 15, FIRST, r2, '=IF(OR($B{r}="",M{r}=0),"",N{r}/M{r})', F_DEC)
        B(ws, 16, FIRST, r2, '=IF($B{r}="","",IFERROR(AVERAGEIFS(RL_Days,RL_Prod,$B{r}),""))', F_DEC)
        B(ws, 17, FIRST, r2, '=IF(OR($B{r}="",M{r}=0),"",INDEX(RL_TS,MATCH($B{r}&"|"&M{r},RL_Key,0)))', F_STAMP)
        B(ws, 18, FIRST, r2, '=IF(OR(L{r}="",N(L{r})=0),"",N{r}/L{r})', F_DEC)
        self.band(ws, 2, 2, 18, FIRST, r2)
        self.status_colors(ws, f"H{FIRST}:H{r2}", {
            "ACTIVE": (OK_BG, OK_FG), "NOT PLACED": (NONE_BG, NONE_FG), "ENDED": (PANEL, NAVY)})
        self.status_colors(ws, f"E{FIRST}:E{r2}", {
            "Snack": ("FDE8D8", "9A4A12"), "Beverage": ("D9ECFB", "0E5A8A")})
        self.table(ws, "tblPilotProducts", 2, 18, r2)
        ws.freeze_panes = f"C{FIRST}"
        for nm, L in [("PP_Name", "B"), ("PP_Type", "E"), ("PP_Flag", "H"), ("PP_Units", "N")]:
            self.name(nm, rng(S_PP, L, FIRST, r2))
        self.protect(ws)

    # ===================================================================== #
    # RESTOCK LOG
    # ===================================================================== #
    def build_restock(self):
        ws = self.sheet(S_RL)
        heads = ["Refill Date", "Refill Time", "Product", "Type", "Qty Before", "Qty Added",
                 "Qty After", "Refilled By", "Previous Refill", "Days Since Prior Refill",
                 "Hours Since Prior Refill", "Operational Notes", "Pilot Day", "Event ID",
                 "Key", "Seq", "Timestamp"]
        last = self.widths(ws, [13, 11, 32, 10, 9, 9, 9, 15, 17, 10, 10, 34, 8, 10, 5, 5, 5],
                           hidden=("P", "Q", "R"))
        self.chrome(ws, S_RL, "One row per refill — type only the BLUE columns. Log the very first "
                    "placement too (Qty Before = 0). Previous refill, cadence and how long the pantry "
                    "lasted are calculated per product. Add new events at the bottom.",
                    last, "PANTRY PILOT")
        kinds = {h: "in" for h in ["Refill Date", "Refill Time", "Product", "Qty Before",
                                   "Qty Added", "Refilled By", "Operational Notes"]}
        self.headers(ws, heads, kinds, height=40)
        r2 = HDR + CAP_REFILLS
        f0 = str(FIRST)
        B = self.body
        B(ws, 2, FIRST, r2, None, F_DATE, kind="in")
        B(ws, 3, FIRST, r2, None, F_TIME, kind="in")
        B(ws, 4, FIRST, r2, None, kind="in", al=LEFT)
        B(ws, 5, FIRST, r2, '=IF(D{r}="","",IFERROR(INDEX(Pil_Type,MATCH(D{r},Pil_Label,0)),"not in pilot"))')
        B(ws, 6, FIRST, r2, None, F_INT, kind="in")
        B(ws, 7, FIRST, r2, None, F_INT, kind="in")
        B(ws, 8, FIRST, r2, '=IF(G{r}="","",N(F{r})+G{r})', F_INT)
        B(ws, 9, FIRST, r2, None, kind="in", al=LEFT)
        B(ws, 10, FIRST, r2, '=IF(OR(Q{r}="",N(Q{r})<=1),"",INDEX($R$' + f0 + ':$R$' + str(r2)
          + ',MATCH(D{r}&"|"&(Q{r}-1),$P$' + f0 + ':$P$' + str(r2) + ',0)))', F_STAMP)
        B(ws, 11, FIRST, r2, '=IF(OR(J{r}="",R{r}=""),"",R{r}-J{r})', F_DEC)
        B(ws, 12, FIRST, r2, '=IF(OR(J{r}="",R{r}=""),"",(R{r}-J{r})*24)', F_DEC)
        B(ws, 13, FIRST, r2, None, kind="in", wrap=True)
        B(ws, 14, FIRST, r2, '=IF(B{r}="","",IF(OR(B{r}<cfg_PilotStart,B{r}>cfg_PilotEnd),"Outside",'
          'B{r}-cfg_PilotStart+1))', "0")
        B(ws, 15, FIRST, r2, '=IF(OR(B{r}="",D{r}=""),"","RF-"&TEXT(ROW()-' + str(HDR) + ',"0000"))')
        B(ws, 16, FIRST, r2, '=IF(Q{r}="","",D{r}&"|"&Q{r})')
        B(ws, 17, FIRST, r2, '=IF(OR(D{r}="",G{r}=""),"",COUNTIFS($D$' + f0 + ':D{r},D{r},$G$' + f0
          + ':G{r},"<>"))')
        B(ws, 18, FIRST, r2, '=IF(B{r}="","",B{r}+N(C{r}))', F_STAMP)
        for a, b in [(5, 5), (8, 8), (10, 12), (14, 15)]:
            self.band(ws, 4, a, b, FIRST, r2)
        self.status_colors(ws, f"E{FIRST}:E{r2}", {
            "Snack": ("FDE8D8", "9A4A12"), "Beverage": ("D9ECFB", "0E5A8A"),
            "not in pilot": (CRIT_BG, CRIT_FG)})
        ws.conditional_formatting.add(f"F{FIRST}:F{r2}", FormulaRule(
            formula=[f'AND($F{FIRST}<>"",$F{FIRST}=0,ISNUMBER($K{FIRST}))'],
            font=Font(bold=True, color=C_RED)))
        self.dv(ws, f"B{FIRST}:B{r2}", "date", "DATE(2020,1,1)", "DATE(2100,12,31)",
                prompt="Date of the refill.", title="Refill date")
        self.dv(ws, f"C{FIRST}:C{r2}", "time", "0", "0.999988",
                prompt="Time of the refill, e.g. 9:30 AM.", title="Refill time")
        self.dv(ws, f"D{FIRST}:D{r2}", "list", "=PilotList",
                prompt="Pick the pantry product.", title="Product")
        self.dv(ws, f"F{FIRST}:G{r2}", "whole", "0", "100000",
                prompt="Units on the shelf before / units added.", title="Quantity")
        self.dv(ws, f"I{FIRST}:I{r2}", "list", "=StaffList", strict=False)
        self.table(ws, "tblRestock", 2, 18, r2, row_height=30)
        ws.freeze_panes = f"A{FIRST}"
        for nm, L in [("RL_Date", "B"), ("RL_Prod", "D"), ("RL_Type", "E"), ("RL_Before", "F"),
                      ("RL_Added", "G"), ("RL_After", "H"), ("RL_Days", "K"), ("RL_Hours", "L"),
                      ("RL_Day", "N"), ("RL_Key", "P"), ("RL_TS", "R")]:
            self.name(nm, rng(S_RL, L, FIRST, r2))
        if self.demo:
            for i, (d, t, prod, before, added, who, note) in enumerate(self.demo["refills"]):
                r = FIRST + i
                for c, v in zip([2, 3, 4, 6, 7, 9, 13], [d, t, prod, before, added, who, note or None]):
                    ws.cell(r, c, v)
        self.protect(ws)

    # ===================================================================== #
    # FEEDBACK LOG
    # ===================================================================== #
    def build_feedback(self):
        ws = self.sheet(S_FB)
        heads = ["Date", "Time", "Feedback Type", "Rating (1-5)", "Product", "Comment",
                 "Area / Team", "Follow-Up Needed", "Owner", "Status", "Resolved Date",
                 "Response Days", "Pilot Day", "Feedback ID"]
        last = self.widths(ws, [13, 10, 13, 9, 26, 40, 16, 10, 15, 11, 13, 10, 8, 10])
        self.chrome(ws, S_FB, "Structured employee feedback — type only the BLUE columns (dropdowns "
                    "wherever possible). Response Days = days to resolution, or days open so far for "
                    "unresolved follow-ups. IDs and pilot day are automatic.", last, "PANTRY PILOT")
        kinds = {h: "in" for h in heads[:11]}
        self.headers(ws, heads, kinds, height=40)
        r2 = HDR + CAP_FEEDBACK
        B = self.body
        B(ws, 2, FIRST, r2, None, F_DATE, kind="in")
        B(ws, 3, FIRST, r2, None, F_TIME, kind="in")
        B(ws, 4, FIRST, r2, None, kind="in")
        B(ws, 5, FIRST, r2, None, "0", kind="in")
        B(ws, 6, FIRST, r2, None, kind="in", al=LEFT)
        B(ws, 7, FIRST, r2, None, kind="in", wrap=True)
        B(ws, 8, FIRST, r2, None, kind="in", al=LEFT)
        B(ws, 9, FIRST, r2, None, kind="in")
        B(ws, 10, FIRST, r2, None, kind="in", al=LEFT)
        B(ws, 11, FIRST, r2, None, kind="in")
        B(ws, 12, FIRST, r2, None, F_DATE, kind="in")
        B(ws, 13, FIRST, r2, '=IF(B{r}="","",IF(L{r}<>"",L{r}-B{r},IF(I{r}="Yes",MAX(0,cfg_AsOf-B{r}),"")))', "0")
        B(ws, 14, FIRST, r2, '=IF(B{r}="","",IF(OR(B{r}<cfg_PilotStart,B{r}>cfg_PilotEnd),"Outside",'
          'B{r}-cfg_PilotStart+1))', "0")
        B(ws, 15, FIRST, r2, '=IF(B{r}="","","FB-"&TEXT(ROW()-' + str(HDR) + ',"0000"))')
        self.band(ws, 2, 13, 15, FIRST, r2)
        self.status_colors(ws, f"D{FIRST}:D{r2}", {
            "Positive": (OK_BG, OK_FG), "Neutral": (NONE_BG, NONE_FG),
            "Concern": (CRIT_BG, CRIT_FG), "Suggestion": ("E6E1FA", "4A3A9E")})
        self.status_colors(ws, f"K{FIRST}:K{r2}", {
            "Open": (CRIT_BG, CRIT_FG), "In Review": (LOW_BG, LOW_FG), "Closed": (OK_BG, OK_FG)})
        self.dv(ws, f"B{FIRST}:B{r2}", "date", "DATE(2020,1,1)", "DATE(2100,12,31)")
        self.dv(ws, f"L{FIRST}:L{r2}", "date", "DATE(2020,1,1)", "DATE(2100,12,31)",
                prompt="Date the follow-up was closed.", title="Resolved date")
        self.dv(ws, f"C{FIRST}:C{r2}", "time", "0", "0.999988")
        self.dv(ws, f"D{FIRST}:D{r2}", "list", '"Positive,Neutral,Concern,Suggestion"')
        self.dv(ws, f"E{FIRST}:E{r2}", "list", '"1,2,3,4,5"',
                prompt="1 = very poor … 5 = excellent", title="Rating")
        self.dv(ws, f"F{FIRST}:F{r2}", "list", "=FeedbackProducts")
        self.dv(ws, f"H{FIRST}:H{r2}", "list", "=TeamList", strict=False)
        self.dv(ws, f"I{FIRST}:I{r2}", "list", '"Yes,No"')
        self.dv(ws, f"J{FIRST}:J{r2}", "list", "=StaffList", strict=False)
        self.dv(ws, f"K{FIRST}:K{r2}", "list", '"Open,In Review,Closed"')
        self.table(ws, "tblFeedback", 2, 15, r2, row_height=30)
        ws.freeze_panes = f"A{FIRST}"
        for nm, L in [("FB_Date", "B"), ("FB_Type", "D"), ("FB_Rating", "E"), ("FB_Prod", "F"),
                      ("FB_FollowUp", "I"), ("FB_Status", "K"), ("FB_Resp", "M"), ("FB_Day", "N")]:
            self.name(nm, rng(S_FB, L, FIRST, r2))
        if self.demo:
            for i, row in enumerate(self.demo["feedback"]):
                for c, v in zip(range(2, 13), row):
                    ws.cell(FIRST + i, c, v)
        self.protect(ws)

    # ===================================================================== #
    # PILOT DASHBOARD
    # ===================================================================== #
    def build_pilot_dash(self):
        ws = self.sheet(S_PD)
        last = self.widths(ws, [10.5] * 16)
        self.chrome(ws, S_PD, "Leadership view of the pantry pilot: consumption, refill cadence and "
                    "employee sentiment, with a rule-based recommendation (rules on the Control Panel).",
                    last, "PILOT DASHBOARD")
        ws.freeze_panes = "A4"
        day = "IF(cfg_AsOf<cfg_PilotStart,0,MIN(cfg_AsOf-cfg_PilotStart+1,cfg_PilotDays))"
        # status banner + progress
        ws.row_dimensions[5].height = 26
        merge(ws, 5, 2, 5, 9,
              f'=IF(cfg_AsOf<cfg_PilotStart,"PILOT STARTS IN "&(cfg_PilotStart-cfg_AsOf)&" DAYS",'
              f'IF(cfg_AsOf>cfg_PilotEnd,"PILOT COMPLETE","PILOT IN PROGRESS  ·  DAY "&{day}&" OF "&cfg_PilotDays))',
              f=font(12, True, NAVY), al=Alignment(vertical="center"))
        merge(ws, 5, 10, 5, 17,
              f'=REPT("█",ROUND({day}/cfg_PilotDays*28,0))&REPT("░",28-ROUND({day}/cfg_PilotDays*28,0))'
              f'&"  "&FIXED({day}/cfg_PilotDays*100,0)&"%"',
              f=font(11, False, C_TEAL), al=Alignment(horizontal="right", vertical="center"))
        ws.row_dimensions[6].height = 16
        merge(ws, 6, 2, 6, 17,
              '="Pilot window: "&cfg_Site&"  ·  "&cfg_PilotDays&" days  ·  refills and feedback logged: "'
              '&COUNT(RL_Added)&" / "&COUNT(FB_Date)',
              f=font(9, False, SLATE, italic=True), al=Alignment(vertical="center"))

        entries = "COUNT(FB_Date)"
        cards = [
            ("PILOT DAY", f"={day}", F_INT, C_NAVY, '="of "&cfg_PilotDays&" planned days"'),
            ("DAYS REMAINING", "=IF(cfg_AsOf<cfg_PilotStart,cfg_PilotDays,MAX(0,cfg_PilotEnd-cfg_AsOf))",
             F_INT, C_PURPLE, '="Ends on day "&cfg_PilotDays'),
            ("REFILL EVENTS", "=COUNT(RL_Added)", F_INT, C_TEAL,
             '=COUNTIF(PP_Flag,"ACTIVE")&" products active"'),
            ("UNITS PLACED", "=SUM(RL_Added)", F_INT, C_BLUE,
             '=IF(' + day + '=0,"—",FIXED(SUM(RL_Added)/' + day + ',1)&" units / day")'),
            ("AVG DAYS BETWEEN REFILLS", '=IFERROR(AVERAGE(RL_Days),"—")', F_DEC, C_NAVY,
             '=IFERROR("≈ "&FIXED(AVERAGE(RL_Hours),0)&" hours per product","No repeat refills yet")'),
            ("FEEDBACK ENTRIES", f"={entries}", F_INT, C_TEAL,
             '=COUNTIF(FB_Type,"Positive")&" positive · "&COUNTIF(FB_Type,"Concern")&" concerns"'),
            ("AVERAGE RATING", '=IFERROR(AVERAGE(FB_Rating),"—")', "0.0", C_GREEN,
             '="out of 5 · "&COUNT(FB_Rating)&" ratings"'),
            ("OPEN FOLLOW-UPS", '=COUNTIFS(FB_FollowUp,"Yes",FB_Status,"<>Closed")', F_INT, C_RED,
             '=IFERROR("Avg response "&FIXED(AVERAGEIFS(FB_Resp,FB_Status,"Closed"),1)&" days","No closed items yet")'),
        ]
        for i, (lab, f, nf, colr, cap) in enumerate(cards):
            r = 8 if i < 4 else 12
            c = 2 + (i % 4) * 4
            self.card(ws, r, c, c + 3, lab, f, nf, colr, cap)
        ws.row_dimensions[11].height = 10

        # decision panel
        ws.row_dimensions[15].height = 12
        self.section(ws, 16, 2, 17, "LEADERSHIP DECISION SIGNAL")
        pos = f'IF({entries}=0,0,COUNTIF(FB_Type,"Positive")/{entries})'
        avg = "IFERROR(AVERAGE(FB_Rating),0)"
        rec = (f'=IF({entries}<cfg_MinFeedback,"GATHER MORE DATA",IF(AND({avg}>=cfg_ExpandRating,'
               f'{pos}>=cfg_ExpandPositive),"EXPAND",IF({avg}>=cfg_ContinueRating,"CONTINUE","ADJUST")))')
        for rr in (17, 18, 19):
            ws.row_dimensions[rr].height = 22
        merge(ws, 17, 2, 19, 5, rec, f=font(18, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        self.name("Pilot_Rec", cell_ref(S_PD, "B17"))
        for text, bg in [("EXPAND", C_GREEN), ("CONTINUE", C_TEAL), ("ADJUST", C_AMBER),
                         ("GATHER MORE DATA", MUTED)]:
            ws.conditional_formatting.add("B17:E19", CellIsRule(
                operator="equal", formula=[f'"{text}"'], fill=fill(bg),
                font=Font(bold=True, color="FFFFFF"), stopIfTrue=True))
        signals = [
            ("Positive feedback share", f'={pos}', F_PCT),
            ("Concerns logged", '=COUNTIF(FB_Type,"Concern")', F_INT),
            ("Refills that found the shelf empty",
             '=COUNTIFS(RL_Before,0,RL_Days,">=0")', F_INT),
            ("Projected units for the full pilot",
             f'=IF({day}=0,"—",SUM(RL_Added)/{day}*cfg_PilotDays)', F_INT),
            ("Snack units  ·  Beverage units",
             '=FIXED(SUMIFS(RL_Added,RL_Type,"Snack"),0)&"  ·  "&FIXED(SUMIFS(RL_Added,RL_Type,"Beverage"),0)', None),
            ("Top product by units", f'=IF(MAX(Pil_Score)<=0,"—",INDEX(Pil_Label,MATCH(MAX(Pil_Score),Pil_Score,0)))', None),
        ]
        for i, (lab, f, nf) in enumerate(signals):
            r = 17 + i % 3
            c1 = 6 if i < 3 else 12
            merge(ws, r, c1, r, c1 + 3, lab, f=font(9.5, False, SLATE), al=LEFT, bd=BOTTOM)
            merge(ws, r, c1 + 4, r, c1 + 5, f, f=font(10.5, True, NAVY), al=RIGHT, bd=BOTTOM, nf=nf)
        merge(ws, 20, 2, 20, 17,
              '="Rules: EXPAND when avg rating ≥ "&FIXED(cfg_ExpandRating,1)&" and positive ≥ "&'
              'FIXED(cfg_ExpandPositive*100,0)&"%  ·  CONTINUE when avg rating ≥ "&FIXED(cfg_ContinueRating,1)&'
              '"  ·  otherwise ADJUST  ·  needs at least "&cfg_MinFeedback&" feedback entries."',
              f=font(8.5, False, MUTED, italic=True), al=Alignment(vertical="center"))

        # chart headers
        self.section(ws, 22, 2, 10, f"UNITS PLACED BY PRODUCT  (top {TOP_N})")
        self.section(ws, 22, 12, 17, "FEEDBACK RATING DISTRIBUTION")

        # data tables below charts
        tr = 43
        self.section(ws, tr, 2, 10, "UNITS PLACED BY PRODUCT — DATA")
        for c1, c2, h in [(2, 2, "#"), (3, 6, "Product"), (7, 8, "Type"), (9, 10, "Units")]:
            merge(ws, tr + 1, c1, tr + 1, c2, h, f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        for k in range(1, TOP_N + 1):
            r = tr + 1 + k
            style(ws.cell(r, 2, k), f=font(9, False, MUTED), al=CENTER, bd=BOTTOM)
            idx = f"MATCH(LARGE(Pil_Score,B{r}),Pil_Score,0)"
            merge(ws, r, 3, r, 6, f'=IF(LARGE(Pil_Score,B{r})<0,"",INDEX(Pil_Label,{idx}))',
                  f=font(10), al=LEFT, bd=BOTTOM)
            merge(ws, r, 7, r, 8, f'=IF(C{r}="","",INDEX(Pil_Type,{idx}))', f=font(10, False, SLATE),
                  al=CENTER, bd=BOTTOM)
            merge(ws, r, 9, r, 10, f'=IF(C{r}="","",INDEX(Pil_Units,{idx}))', f=font(10, True, NAVY),
                  al=CENTER, bd=BOTTOM, nf=F_INT)

        self.section(ws, tr, 12, 17, "RATING DISTRIBUTION — DATA")
        for c1, c2, h in [(12, 13, "Rating"), (14, 15, "Responses"), (16, 17, "Share")]:
            merge(ws, tr + 1, c1, tr + 1, c2, h, f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        rating_colors = [C_RED, "E8663A", C_AMBER, "7DBB5A", C_GREEN]
        for i in range(5):
            r = tr + 2 + i
            merge(ws, r, 12, r, 13, f"{i + 1} ★", f=font(10, True, rating_colors[i]), al=CENTER, bd=BOTTOM)
            merge(ws, r, 14, r, 15, f"=COUNTIF(FB_Rating,{i + 1})", f=font(10), al=CENTER, bd=BOTTOM, nf=F_INT)
            merge(ws, r, 16, r, 17, f'=IF(COUNT(FB_Rating)=0,"",N{r}/COUNT(FB_Rating))', f=font(10),
                  al=CENTER, bd=BOTTOM, nf=F_PCT)

        # weekly trend table
        weeks = math.ceil(self.pilot_days / 7)
        wr = tr + TOP_N + 4
        self.section(ws, wr, 2, 17, "WEEK-BY-WEEK TREND")
        for c1, c2, h in [(2, 3, "Pilot Week"), (4, 5, "Days"), (6, 7, "Refill Events"),
                          (8, 9, "Units Placed"), (10, 11, "Feedback"), (12, 13, "Avg Rating"),
                          (14, 15, "Concerns"), (16, 17, "Empty on Arrival")]:
            merge(ws, wr + 1, c1, wr + 1, c2, h, f=font(9, True, "FFFFFF"), fl=fill(NAVY), al=CENTER)
        for w in range(1, weeks + 1):
            r = wr + 1 + w
            lo, hi = (w - 1) * 7 + 1, w * 7
            dr = f'RL_Day,">={lo}",RL_Day,"<={hi}"'
            fr = f'FB_Day,">={lo}",FB_Day,"<={hi}"'
            vals = [
                (2, 3, f"Week {w}", None),
                (4, 5, f'="{lo}–"&MIN({hi},cfg_PilotDays)', None),
                (6, 7, f"=COUNTIFS({dr})", F_INT),
                (8, 9, f"=SUMIFS(RL_Added,{dr})", F_INT),
                (10, 11, f"=COUNTIFS({fr})", F_INT),
                (12, 13, f'=IFERROR(AVERAGEIFS(FB_Rating,{fr}),"—")', "0.0"),
                (14, 15, f'=COUNTIFS({fr},FB_Type,"Concern")', F_INT),
                (16, 17, f'=COUNTIFS({dr},RL_Before,0,RL_Days,">=0")', F_INT),
            ]
            for c1, c2, v, nf in vals:
                merge(ws, r, c1, r, c2, v, f=font(10, c1 == 2, NAVY if c1 == 2 else "1F2A44"),
                      al=CENTER, bd=BOTTOM, nf=nf)
            if w % 2 == 0:
                for c in range(2, 18):
                    ws.cell(r, c).fill = fill(BAND)

        # charts
        b = self.hbar(ws, ref_vals=(9, tr + 2, tr + 1 + TOP_N), ref_cats=(3, tr + 2, tr + 1 + TOP_N),
                      color=C_TEAL)
        b.width, b.height = 19.5, 9.4
        ws.add_chart(b, "B23")
        d = DoughnutChart(holeSize=58)
        d.add_data(Reference(ws, min_col=14, min_row=tr + 2, max_row=tr + 6), titles_from_data=False)
        d.set_categories(Reference(ws, min_col=12, min_row=tr + 2, max_row=tr + 6))
        s = d.series[0]
        for i, colr in enumerate(rating_colors):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = colr
            pt.graphicalProperties.line.solidFill = "FFFFFF"
            s.dPt.append(pt)
        s.dLbls = DataLabelList(); s.dLbls.showVal = True
        for a in ("showSerName", "showCatName", "showLegendKey", "showPercent", "showLeaderLines"):
            setattr(s.dLbls, a, False)
        d.legend.position = "r"
        d.width, d.height = 11.5, 9.4
        ws.add_chart(d, "L23")
        self.protect(ws)

    # ===================================================================== #
    # HOME (launcher)
    # ===================================================================== #
    def build_home(self):
        ws = self.sheet(S_HOME, first=True)
        last = self.widths(ws, [10.5] * 16)
        self.chrome(ws, S_HOME, "Welcome. Pick a module below. Only three things are ever typed by "
                    "hand — counts, refills and feedback — everything else runs itself.",
                    last, "HOME")
        ws.freeze_panes = "A4"
        ws.row_dimensions[5].height = 10
        if self.demo:
            merge(ws, 6, 2, 6, 17, "DEMO DATA — sample products and entries for exploring the "
                  "system. Build from the real source file before going live.",
                  f=font(10, True, CRIT_FG), fl=fill(CRIT_BG), al=CENTER)
            ws.row_dimensions[6].height = 24
        else:
            merge(ws, 6, 2, 6, 17, '=IF(SrcCount=0,"NO PRODUCTS LOADED — open Source Data and paste '
                  'Category, Item Name, Item #, Recent Date Received and PAR Value from the source file.","")',
                  f=font(10, True, CRIT_FG), al=CENTER)
            ws.conditional_formatting.add("B6:Q6", FormulaRule(formula=['$B$6<>""'], fill=fill(CRIT_BG)))
            ws.row_dimensions[6].height = 24

        # live status strip
        self.section(ws, 7, 2, 17, "RIGHT NOW")
        live = [
            ("CRITICAL ITEMS", '=COUNTIF(MS_Status,"CRITICAL")', F_INT, C_RED),
            ("ITEMS TO ORDER", '=COUNTIF(OL_Flag,"ORDER")', F_INT, C_BLUE),
            ("PILOT DAY", "=IF(cfg_AsOf<cfg_PilotStart,0,MIN(cfg_AsOf-cfg_PilotStart+1,cfg_PilotDays))",
             F_INT, C_TEAL),
            ("PILOT SIGNAL", "=Pilot_Rec", None, C_NAVY),
        ]
        for i, (lab, f, nf, colr) in enumerate(live):
            c = 2 + i * 4
            self.card(ws, 8, c, c + 3, lab, f, nf, colr, "")
        ws.cell(9, 14).font = font(15, True, "FFFFFF")

        def tile(r, c, title, desc, sheet, colr, entry=False):
            for rr in (r, r + 1, r + 2):
                ws.row_dimensions[rr].height = 20 if rr != r + 1 else 32
            t = merge(ws, r, c, r, c + 3, ("✎  " if entry else "") + title, f=font(11, True, "FFFFFF"),
                      fl=fill(colr), al=Alignment(vertical="bottom", indent=1))
            link(t, sheet, f"Open {sheet}")
            dsc = merge(ws, r + 1, c, r + 2, c + 3, desc, f=font(8.5, False, "E8EEFA"), fl=fill(colr),
                        al=Alignment(vertical="top", wrap_text=True, indent=1))
            link(dsc, sheet, f"Open {sheet}")

        self.section(ws, 13, 2, 17, "SYSTEM 1  ·  INVENTORY MANAGEMENT")
        s1 = [
            ("CONTROL PANEL", "Thresholds, lead time, order month, pilot dates, lists.", S_CP, "5B6B8C", False),
            ("INVENTORY MASTER", "Every product with stock %, status and suggested order.", S_MS, NAVY, False),
            ("WEEKLY COUNT", "Log the weekly physical count here.", S_CL, IN_HDR, True),
            ("ORDER LIST", "Auto-built purchase list with priority and need-by date.", S_OL, NAVY, False),
        ]
        for i, (a, b, s, colr, e) in enumerate(s1):
            tile(14, 2 + i * 4, a, b, s, colr, e)
        ws.row_dimensions[17].height = 8
        tile(18, 2, "INVENTORY DASHBOARD", "KPIs, status mix and order units by category.", S_ID, INK)
        tile(18, 6, "SOURCE DATA", "Product master (swap this to launch another office).", S_SRC, MUTED)

        self.section(ws, 22, 2, 17, "SYSTEM 2  ·  60-DAY PANTRY PILOT" if self.pilot_days == 60
                     else "SYSTEM 2  ·  PANTRY PILOT")
        s2 = [
            ("PILOT PRODUCTS", "Snacks & beverages in the pantry with usage stats.", S_PP, C_TEAL, False),
            ("RESTOCK LOG", "Log every refill (date, time, qty).", S_RL, IN_HDR, True),
            ("FEEDBACK LOG", "Log employee feedback and follow-ups.", S_FB, IN_HDR, True),
            ("PILOT DASHBOARD", "Leadership view and continue / expand / adjust signal.", S_PD, "0B5E66", False),
        ]
        for i, (a, b, s, colr, e) in enumerate(s2):
            tile(23, 2 + i * 4, a, b, s, colr, e)

        self.section(ws, 27, 2, 17, "YOUR ROUTINE  (a few entries per week)")
        steps = [
            ("1", "Every week", "Weekly Count: one row per product counted — date, product, quantity, "
                                "your name. Status, usage and the order list update automatically."),
            ("2", "Every refill", "Restock Log: date, time, product, qty before, qty added, your name. "
                                  "Cadence and hours-lasted calculate themselves."),
            ("3", "Any feedback", "Feedback Log: date, type, rating, product, comment, team, follow-up. "
                                  "Close items by setting Status = Closed and the resolved date."),
            ("4", "Monthly", "Open Order List, filter Order Flag = ORDER, place the order. "
                             "Change the order month on the Control Panel."),
        ]
        for i, (n, a, b) in enumerate(steps):
            r = 28 + i
            ws.row_dimensions[r].height = 30
            style(ws.cell(r, 2, n), f=font(14, True, "FFFFFF"), fl=fill(IN_HDR if i < 3 else NAVY), al=CENTER)
            merge(ws, r, 3, r, 4, a, f=font(10, True, NAVY), al=LEFT, bd=BOTTOM)
            merge(ws, r, 5, r, 17, b, f=font(9.5, False, SLATE), al=LEFT, bd=BOTTOM)
        r = 33
        style(ws.cell(r, 2, ""), fl=fill(IN_FILL), bd=BOX)
        merge(ws, r, 3, r, 7, "Blue = your entries", f=font(9, True, IN_FONT), al=LEFT)
        style(ws.cell(r, 8, ""), fl=fill(SET_FILL), bd=BOX)
        merge(ws, r, 9, r, 12, "Yellow = one-time setup", f=font(9, True, SET_FONT), al=LEFT)
        style(ws.cell(r, 13, ""), fl=fill(PANEL), bd=BOX)
        merge(ws, r, 14, r, 17, "White / grey = automatic", f=font(9, True, NAVY), al=LEFT)
        self.protect(ws)

    # ===================================================================== #
    def build(self):
        self.build_home()
        self.build_control()
        self.build_source()
        self.build_master()
        self.build_count()
        self.build_orders()
        self.build_inv_dash()
        self.build_pilot_products()
        self.build_restock()
        self.build_feedback()
        self.build_pilot_dash()
        self.build_engine()
        wb = self.wb
        order = [S_HOME, S_CP, S_MS, S_CL, S_OL, S_ID, S_PP, S_RL, S_FB, S_PD, S_SRC, S_EN]
        wb._sheets = [self.ws[n] for n in order]
        wb.active = 0
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = ws.title == S_HOME
            # print-ready: landscape, one page wide, header rows repeated
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_options.horizontalCentered = True
            ws.page_margins.left = ws.page_margins.right = 0.4
            if ws.tables:
                ws.print_title_rows = f"1:{HDR}"
        for n, rows in [(S_HOME, 34), (S_ID, 76), (S_PD, 70), (S_CP, 45)]:
            self.ws[n].print_area = f"A1:{col_letter(self.last_col[n])}{rows}"
        wb.calculation.fullCalcOnLoad = True
        wb.properties.title = f"{self.site} Operations Hub"
        wb.properties.creator = "The Automation Lab"
        return wb


# --------------------------------------------------------------------------- #
def parse_date(s):
    return dt.datetime.strptime(s, "%Y-%m-%d").date()


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", type=Path, help="Source product workbook (.xlsx)")
    ap.add_argument("--site", default="SJUB", help="Site / office name (default SJUB)")
    ap.add_argument("--pilot-start", type=parse_date, default=dt.date(2026, 9, 28),
                    help="Pilot start date YYYY-MM-DD (default 2026-09-28)")
    ap.add_argument("--pilot-days", type=int, default=60, help="Pilot length in days (default 60)")
    ap.add_argument("--order-month", type=parse_date, default=None,
                    help="Fixed order month YYYY-MM-01 (default: follows the report date)")
    ap.add_argument("--snack-cats", default=None,
                    help="Comma-separated source categories that are Snack (default: auto-detect)")
    ap.add_argument("--beverage-cats", default=None,
                    help="Comma-separated source categories that are Beverage (default: auto-detect)")
    ap.add_argument("--demo", action="store_true", help="Fill with clearly-labelled sample data")
    ap.add_argument("--xltx", action="store_true", help="Save as an Excel template (.xltx)")
    ap.add_argument("--out", type=Path, default=None, help="Output file")
    a = ap.parse_args(argv)

    demo = None
    as_of = None
    if a.demo:
        products = demo_products()
    elif a.source:
        products = load_source(a.source)
    else:
        products = []
    snack, bev = (["Snacks"], ["Beverages"]) if a.demo else guess_pilot_categories(products)
    if a.snack_cats is not None:
        snack = [c.strip() for c in a.snack_cats.split(",") if c.strip()]
    if a.beverage_cats is not None:
        bev = [c.strip() for c in a.beverage_cats.split(",") if c.strip()]
    if a.demo:
        counts, refills, fb, staff, teams = demo_entries(products, a.pilot_start, snack, bev)
        demo = {"counts": counts, "refills": refills, "feedback": fb, "staff": staff, "teams": teams}
        as_of = a.pilot_start + dt.timedelta(days=23)
    if len(products) > CAP_PRODUCTS:
        sys.exit(f"{len(products)} products exceeds capacity {CAP_PRODUCTS}; raise CAP_PRODUCTS.")

    order_month = a.order_month
    hub = Hub(a.site, products, a.pilot_start, a.pilot_days, order_month, snack, bev,
              as_of=as_of, demo=demo)
    wb = hub.build()
    out = a.out or Path(f"{a.site}_Operations_Hub{'_DEMO' if a.demo else ''}"
                        f"{'.xltx' if a.xltx else '.xlsx'}")
    out.parent.mkdir(parents=True, exist_ok=True)
    if a.xltx:
        wb.template = True
    wb.save(out)
    print(f"Saved {out}  ({len(products)} products · snack cats {snack} · beverage cats {bev})")


if __name__ == "__main__":
    main()
